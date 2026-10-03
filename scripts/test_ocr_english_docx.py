"""English OCR fallback contracts without an OCR engine or HTTP server."""

from __future__ import annotations

from pathlib import Path
import contextlib
import importlib.util
import io
import tempfile
from types import SimpleNamespace
import unittest
from unittest.mock import patch

import pymupdf as fitz
from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION_START

from ocr_english_docx import (
    EmptyOcrPageError,
    OcrRuntimeUnavailableError,
    recognize_ocr_transcripts,
    supports_latin_ocr,
    validate_page_transcripts,
    write_english_ocr_docx,
)

SPEC = importlib.util.spec_from_file_location("pdf_docx_english_branch", Path(__file__).with_name("pdf-to-docx.py"))
assert SPEC is not None and SPEC.loader is not None
converter = importlib.util.module_from_spec(SPEC)
SPEC.loader.exec_module(converter)


class EnglishOcrDocxTests(unittest.TestCase):
    def setUp(self) -> None:
        self.directory = tempfile.TemporaryDirectory(prefix="onlymypdf-english-ocr-")
        self.addCleanup(self.directory.cleanup)
        self.source = Path(self.directory.name) / "source.pdf"
        self.output = Path(self.directory.name) / "output.docx"

    def source_pages(self, *sizes: tuple[int, int]) -> None:
        with fitz.open() as document:
            for width, height in sizes:
                document.new_page(width=width, height=height)
            document.save(self.source)

    def test_actual_latin_letters_select_the_fallback_without_language_claims(self) -> None:
        self.assertTrue(supports_latin_ocr(["Invoice 58324\nTotal amount 1475.25"]))
        self.assertTrue(supports_latin_ocr(["Caf\u00e9 reference"]))
        self.assertFalse(supports_latin_ocr(["English reference\n\u041f\u0440\u0438\u0432\u0435\u0442"]))
        self.assertFalse(supports_latin_ocr(["12345.67"]))
        self.assertFalse(supports_latin_ocr([" \n"]))

    def test_words_spaces_and_repeated_lines_remain_editable_without_images(self) -> None:
        self.source_pages((612, 792))
        lines = ["  Account  number:  001234  ", "This remains searchable and", "editable.", "Repeated line", "Repeated line", "Total amount 1475.25"]
        write_english_ocr_docx(str(self.source), str(self.output), ["\n".join(lines)])
        document = Document(self.output)
        self.assertEqual([paragraph.text for paragraph in document.paragraphs], lines)
        self.assertEqual(len(document.tables), 0)
        self.assertEqual(len(document.inline_shapes), 0)
        self.assertEqual(document.styles["Normal"].font.size.pt, 11)
        self.assertGreater((document.sections[0].page_width - document.sections[0].left_margin - document.sections[0].right_margin) / 914400, 7)

    def test_source_sections_keep_orientation_and_last_page_text(self) -> None:
        self.source_pages((612, 792), (842, 595))
        write_english_ocr_docx(str(self.source), str(self.output), ["FIRST PAGE invoice 58324", "LAST PAGE amount 1475.25"])
        document = Document(self.output)
        self.assertEqual(len(document.sections), 2)
        self.assertEqual(document.sections[0].page_width.pt, 612)
        self.assertEqual(document.sections[0].page_height.pt, 792)
        self.assertEqual(document.sections[0].orientation, WD_ORIENT.PORTRAIT)
        self.assertEqual(document.sections[1].page_width.pt, 842)
        self.assertEqual(document.sections[1].page_height.pt, 595)
        self.assertEqual(document.sections[1].orientation, WD_ORIENT.LANDSCAPE)
        self.assertEqual(document.sections[1].start_type, WD_SECTION_START.NEW_PAGE)
        text = [paragraph.text for paragraph in document.paragraphs if paragraph.text]
        self.assertEqual(text, ["FIRST PAGE invoice 58324", "LAST PAGE amount 1475.25"])

    def test_empty_later_page_is_not_hidden_by_readable_first_page(self) -> None:
        self.source_pages((612, 792), (612, 792))
        with self.assertRaisesRegex(EmptyOcrPageError, "page 2"):
            write_english_ocr_docx(str(self.source), str(self.output), ["Readable invoice 58324", " \n\t"])
        self.assertFalse(self.output.exists())

    def test_ocr_fallback_uses_original_geometry_not_scaled_intermediate(self) -> None:
        self.source_pages((612, 792), (792, 612))
        texts = ["Original portrait invoice amount 1475.25", "Original landscape totals reference 71936"]
        diagnostics = io.StringIO()
        with patch.object(converter, "_build_searchable_ocr_pdf", side_effect=self.prepared_ocr(texts)), \
             patch.object(converter, "install_pdf2docx_devanagari_support"), \
             patch("ocr_english_docx.recognize_ocr_transcripts", return_value=texts), \
             patch("pdf2docx.Converter", self.incomplete_converter()), \
             contextlib.redirect_stderr(diagnostics):
            self.assertTrue(converter.convert_scanned_pdf_with_ocr(str(self.source), str(self.output), language="eng"))
        sections = Document(self.output).sections
        self.assertEqual([(section.page_width.pt, section.page_height.pt) for section in sections], [(612, 792), (792, 612)])
        self.assertIn("method=editable-latin-transcript", diagnostics.getvalue())

    def prepared_ocr(self, texts: list[str]):
        def build(_source: str, output: str, **_options):
            with fitz.open() as document:
                for index in range(2):
                    size = (1402.5, 1815) if index == 0 else (1815, 1402.5)
                    document.new_page(width=size[0], height=size[1])
                document.save(output)
            return 2, sum(len(text) for text in texts), texts
        return build

    @staticmethod
    def incomplete_converter():
        class IncompleteConverter:
            def __init__(self, *_args):
                pass

            def convert(self, path: str, **_options):
                document = Document()
                document.add_paragraph("Insufficient")
                document.save(path)

            def close(self):
                pass
        return IncompleteConverter

    def test_readable_first_page_cannot_hide_empty_page_in_actual_ocr_branch(self) -> None:
        self.source_pages((612, 792), (612, 792))
        texts = ["Readable invoice with enough text to pass an aggregate-only gate 58324", ""]
        diagnostics = io.StringIO()
        with patch.object(converter, "_build_searchable_ocr_pdf", side_effect=self.prepared_ocr(texts)), \
             patch.object(converter, "install_pdf2docx_devanagari_support"), \
             patch("ocr_english_docx.recognize_ocr_transcripts", return_value=texts), \
             patch("pdf2docx.Converter") as engine, \
             contextlib.redirect_stderr(diagnostics):
            self.assertFalse(converter.convert_scanned_pdf_with_ocr(str(self.source), str(self.output), language="eng"))
        engine.assert_not_called()
        self.assertIn("WARN OCR text below gate: chars=0 need=1", diagnostics.getvalue())
        self.assertIn("page 2", diagnostics.getvalue())
        self.assertFalse(self.output.exists())

    def test_successful_converter_output_does_not_use_the_new_fallback(self) -> None:
        self.source_pages((612, 792), (612, 792))
        texts = ["First page keeps readable English text", "Second page keeps all expected editable words"]

        class CompleteConverter(self.incomplete_converter()):
            def convert(self, path: str, **_options):
                document = Document()
                for text in texts:
                    document.add_paragraph(text)
                document.save(path)

        with patch.object(converter, "_build_searchable_ocr_pdf", side_effect=self.prepared_ocr(texts)), \
             patch.object(converter, "install_pdf2docx_devanagari_support"), \
             patch("pdf2docx.Converter", CompleteConverter), \
             patch("ocr_english_docx.recognize_ocr_transcripts") as recognition, \
             patch("ocr_english_docx.write_english_ocr_docx") as fallback:
            self.assertTrue(converter.convert_scanned_pdf_with_ocr(str(self.source), str(self.output), language="eng"))
        fallback.assert_not_called()
        recognition.assert_not_called()

    def test_non_latin_ocr_retains_the_existing_reference_fallback(self) -> None:
        self.source_pages((612, 792), (612, 792))
        texts = ["\u039a\u03b1\u03bb\u03b7\u03bc\u03ad\u03c1\u03b1 English mixed reference paragraph", "A second full paragraph for the existing routine"]

        def existing_reference(_source, path, transcripts):
            document = Document()
            for text in transcripts:
                document.add_paragraph(text)
            document.save(path)
            return True

        with patch.object(converter, "_build_searchable_ocr_pdf", side_effect=self.prepared_ocr(texts)), \
             patch.object(converter, "install_pdf2docx_devanagari_support"), \
             patch("pdf2docx.Converter", self.incomplete_converter()), \
             patch.object(converter, "_build_ocr_reference_transcript_docx", side_effect=existing_reference) as existing, \
             patch("ocr_english_docx.recognize_ocr_transcripts") as recognition, \
             patch("ocr_english_docx.write_english_ocr_docx") as fallback:
            self.assertTrue(converter.convert_scanned_pdf_with_ocr(str(self.source), str(self.output), language="eng"))
        existing.assert_called_once()
        fallback.assert_not_called()
        recognition.assert_not_called()

    def test_raw_recognizer_text_keeps_spaces_and_repeated_rows_in_actual_branch(self) -> None:
        self.source_pages((612, 792), (612, 792))
        hidden = ["Each paragraph hasa different purpose.", "Reference only 71936"]
        raw = ["Each paragraph has a different purpose.\nRepeated row 120.25\nRepeated row 120.25\n", "Analysis package 2 120.25 240.50\nLast page reference 71936\n"]
        with patch.object(converter, "_build_searchable_ocr_pdf", side_effect=self.prepared_ocr(hidden)), \
             patch.object(converter, "install_pdf2docx_devanagari_support"), \
             patch("ocr_english_docx.recognize_ocr_transcripts", return_value=raw) as recognition, \
             patch("pdf2docx.Converter", self.incomplete_converter()):
            self.assertTrue(converter.convert_scanned_pdf_with_ocr(str(self.source), str(self.output), language="eng"))
        text = [paragraph.text for paragraph in Document(self.output).paragraphs if paragraph.text]
        self.assertEqual(text, [*raw[0].splitlines(), *raw[1].splitlines()])
        recognition.assert_called_once_with(str(self.source), language="eng", dpi=220)

    def test_empty_hidden_page_can_recover_from_actual_raw_text_before_rejection(self) -> None:
        self.source_pages((612, 792), (792, 612))
        hidden = ["Readable first page with sufficient detail and reference 58324", ""]
        raw = [hidden[0], "Recovered last page amount 1475.25 and reference 71936"]
        with patch.object(converter, "_build_searchable_ocr_pdf", side_effect=self.prepared_ocr(hidden)), \
             patch.object(converter, "install_pdf2docx_devanagari_support"), \
             patch("ocr_english_docx.recognize_ocr_transcripts", return_value=raw) as recognition, \
             patch("pdf2docx.Converter", self.incomplete_converter()):
            self.assertTrue(converter.convert_scanned_pdf_with_ocr(str(self.source), str(self.output), language="eng"))
        text = [paragraph.text for paragraph in Document(self.output).paragraphs if paragraph.text]
        self.assertEqual(text, raw)
        recognition.assert_called_once()

    def test_short_latin_hidden_text_can_recover_before_aggregate_rejection(self) -> None:
        self.source_pages((612, 792), (612, 792))
        hidden = ["Brief", ""]
        raw = ["First recovered page reference 58324", "Second recovered page amount 1475.25"]
        with patch.object(converter, "_build_searchable_ocr_pdf", side_effect=self.prepared_ocr(hidden)), \
             patch.object(converter, "install_pdf2docx_devanagari_support"), \
             patch("ocr_english_docx.recognize_ocr_transcripts", return_value=raw), \
             patch("pdf2docx.Converter", self.incomplete_converter()):
            self.assertTrue(converter.convert_scanned_pdf_with_ocr(str(self.source), str(self.output), language="eng"))
        self.assertEqual([paragraph.text for paragraph in Document(self.output).paragraphs if paragraph.text], raw)

    def test_nonempty_raw_text_below_existing_gate_is_unreadable_without_artifact(self) -> None:
        self.source_pages((612, 792), (612, 792))
        hidden = ["First hidden page reference 58324", "Second hidden page reference 71936"]
        with patch.object(converter, "_build_searchable_ocr_pdf", side_effect=self.prepared_ocr(hidden)), \
             patch.object(converter, "install_pdf2docx_devanagari_support"), \
             patch("ocr_english_docx.recognize_ocr_transcripts", return_value=["Short", "Tiny"]), \
             patch("pdf2docx.Converter", self.incomplete_converter()), \
             contextlib.redirect_stderr(io.StringIO()) as diagnostics:
            self.assertFalse(converter.convert_scanned_pdf_with_ocr(str(self.source), str(self.output), language="eng"))
        self.assertIn("WARN OCR text below gate: chars=9 need=24", diagnostics.getvalue())
        self.assertNotIn("WARN OCR conversion failed", diagnostics.getvalue())
        self.assertFalse(self.output.exists())

    def test_runtime_failure_is_not_reported_as_unreadable_input(self) -> None:
        self.source_pages((612, 792), (612, 792))
        texts = ["Readable first reference 58324", "Readable last reference 71936"]
        for error, expected in [
            (OcrRuntimeUnavailableError("Missing executable"), "ERROR OCR_REQUIRED Tesseract or requested language data unavailable"),
            (RuntimeError("Unexpected engine failure"), "WARN OCR conversion failed: Unexpected engine failure"),
        ]:
            with self.subTest(error=error), \
                 patch.object(converter, "_build_searchable_ocr_pdf", side_effect=self.prepared_ocr(texts)), \
                 patch.object(converter, "install_pdf2docx_devanagari_support"), \
                 patch("ocr_english_docx.recognize_ocr_transcripts", side_effect=error), \
                 patch("pdf2docx.Converter", self.incomplete_converter()), \
                 contextlib.redirect_stderr(io.StringIO()) as diagnostics:
                self.assertFalse(converter.convert_scanned_pdf_with_ocr(str(self.source), str(self.output), language="eng"))
            self.assertIn(expected, diagnostics.getvalue())
            self.assertNotIn("WARN OCR text below gate", diagnostics.getvalue())
            self.assertFalse(self.output.exists())

    def test_cli_uses_raw_psm3_text_and_keeps_word_boundaries(self) -> None:
        self.source_pages((612, 792), (792, 612))
        raw = ["First  repeated line\nFirst  repeated line\n", "Last amount 1475.25\n"]
        calls = []

        def run(arguments, **options):
            Path(arguments[2] + ".txt").write_text(raw[len(calls)], encoding="utf-8")
            calls.append((arguments, options))
            return SimpleNamespace(returncode=0, stderr="")

        with patch("ocr_english_docx.shutil.which", return_value="tesseract-test"), \
             patch("ocr_english_docx.subprocess.run", side_effect=run):
            recognized = recognize_ocr_transcripts(str(self.source), language="eng", dpi=220)
        self.assertEqual(recognized, raw)
        self.assertEqual(len(calls), 2)
        self.assertEqual(calls[0][0][3:], ["-l", "eng", "--dpi", "220", "--psm", "3", "txt"])
        self.assertEqual(calls[0][1]["timeout"], 90)

    def test_missing_second_cli_output_cannot_reuse_first_page(self) -> None:
        self.source_pages((612, 792), (612, 792))
        calls = []

        def run(arguments, **_options):
            if not calls:
                Path(arguments[2] + ".txt").write_text("First page only", encoding="utf-8")
            calls.append(arguments)
            return SimpleNamespace(returncode=0, stderr="")

        with patch("ocr_english_docx.shutil.which", return_value="tesseract-test"), \
             patch("ocr_english_docx.subprocess.run", side_effect=run), \
             self.assertRaisesRegex(RuntimeError, "produced no transcript"):
            recognize_ocr_transcripts(str(self.source), language="eng", dpi=220)

    def test_cli_missing_executable_or_language_data_are_runtime_failures(self) -> None:
        self.source_pages((612, 792))
        with patch("ocr_english_docx.shutil.which", return_value=None), \
             self.assertRaises(OcrRuntimeUnavailableError):
            recognize_ocr_transcripts(str(self.source), language="eng", dpi=220)
        with patch("ocr_english_docx.shutil.which", return_value="tesseract-test"), \
             patch("ocr_english_docx.subprocess.run", return_value=SimpleNamespace(returncode=1, stderr="Failed loading language 'eng'")), \
             self.assertRaises(OcrRuntimeUnavailableError):
            recognize_ocr_transcripts(str(self.source), language="eng", dpi=220)

    def test_cli_unexpected_failure_does_not_become_missing_runtime(self) -> None:
        self.source_pages((612, 792))
        with patch("ocr_english_docx.shutil.which", return_value="tesseract-test"), \
             patch("ocr_english_docx.subprocess.run", return_value=SimpleNamespace(returncode=2, stderr="unexpected engine failure")), \
             self.assertRaisesRegex(RuntimeError, "exit 2") as error:
            recognize_ocr_transcripts(str(self.source), language="eng", dpi=220)
        self.assertNotIsInstance(error.exception, OcrRuntimeUnavailableError)

    def test_missing_page_transcript_is_an_error_not_a_dropped_source_page(self) -> None:
        self.source_pages((612, 792), (612, 792))
        with self.assertRaisesRegex(ValueError, "count does not match"):
            write_english_ocr_docx(str(self.source), str(self.output), ["Only first page exists"])
        self.assertFalse(self.output.exists())

    def test_non_latin_text_cannot_accidentally_enter_the_new_writer(self) -> None:
        self.source_pages((612, 792))
        with self.assertRaisesRegex(ValueError, "Latin-script"):
            write_english_ocr_docx(str(self.source), str(self.output), ["English \u039a\u03b1\u03bb\u03b7\u03bc\u03ad\u03c1\u03b1"])
        self.assertFalse(self.output.exists())

    def test_numbers_only_page_remains_valid_within_an_english_document(self) -> None:
        validate_page_transcripts(["Invoice", "1475.25"], 2)
        self.assertTrue(supports_latin_ocr(["Invoice", "1475.25"]))


if __name__ == "__main__":
    unittest.main()
