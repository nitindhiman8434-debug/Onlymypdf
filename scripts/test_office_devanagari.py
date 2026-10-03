"""Regression tests for editable Hindi clusters and consistent Word table grids."""
import importlib.util
import os
from pathlib import Path
import unittest

from docx import Document
from docx.shared import Inches
from docx.oxml.ns import qn
from pptx import Presentation

from office_devanagari import (
    devanagari_font, install_pdf2docx_devanagari_support, line_text,
    normalize_devanagari_blocks, synchronize_fixed_table_grid,
)


def span(text, box=(10, 20, 50, 35), **extra):
    return dict(text=text, bbox=box, font="NirmalaUI", size=12, color=0, flags=0, **extra)


def line(text, box=(10, 20, 50, 35)):
    return dict(bbox=box, dir=(1, 0), spans=[span(text, box)])


def blocks(*lines):
    return [dict(type=0, bbox=(10, 20, 150, 100), lines=list(lines))]


class ClusterTests(unittest.TestCase):
    def test_combining_fragment_joins_in_logical_not_x_order(self):
        original = blocks(line("कुल राश"), line("ि", (35, 20, 40, 35)))
        normalized = normalize_devanagari_blocks(original)
        self.assertEqual(line_text(normalized[0]["lines"][0]), "कुल राशि")
        self.assertEqual(len(normalized[0]["lines"]), 1)
        self.assertEqual(len(original[0]["lines"]), 2, "do not mutate source data")

    def test_fragment_across_adjacent_blocks_joins(self):
        original = blocks(line("ह")) + blocks(line("िंदी", (9, 20, 70, 35)))
        normalized = normalize_devanagari_blocks(original)
        self.assertEqual(len(normalized), 1)
        self.assertEqual(line_text(normalized[0]["lines"][0]), "हिंदी")

    def test_different_row_is_not_joined(self):
        normalized = normalize_devanagari_blocks(blocks(line("राश"), line("ि", (35, 40, 40, 55))))
        self.assertEqual(len(normalized[0]["lines"]), 2)

    def test_separate_table_cell_is_not_joined(self):
        normalized = normalize_devanagari_blocks(blocks(line("राश"), line("ि", (80, 20, 90, 35))))
        self.assertEqual(len(normalized[0]["lines"]), 2)

    def test_latin_paths_are_unchanged(self):
        original = blocks(line("Revenue"), line("1250.50", (80, 20, 150, 35)))
        self.assertIs(normalize_devanagari_blocks(original), original)

    def test_whitespace_font_runs_survive_and_style_changes_remain(self):
        sample = line("")
        sample["spans"] = [span("कुल"), span("\u00a0"), span("राशि")]
        sample["spans"][1]["font"] = "ArialMT"
        sample["spans"][1]["flags"] = 4
        sample["spans"].append(dict(span(" विशेष"), flags=16))
        normalized = normalize_devanagari_blocks(blocks(sample))[0]["lines"][0]
        self.assertEqual([s["text"] for s in normalized["spans"]], ["कुल राशि", " विशेष"])

    def test_raw_characters_preserve_every_codepoint(self):
        sample = blocks(line("राश"), line("ि", (35, 20, 40, 35)))
        for item in sample[0]["lines"]:
            for run in item["spans"]:
                run["chars"] = [dict(c=c, bbox=run["bbox"], origin=(10, 30)) for c in run.pop("text")]
        normalized = normalize_devanagari_blocks(sample)
        self.assertEqual(line_text(normalized[0]["lines"][0]), "राशि")

    def test_fonts_use_family_names_and_linux_fallback(self):
        self.assertEqual(devanagari_font("ABCDEF+NotoSansDevanagari-Regular"), "Noto Sans Devanagari")
        self.assertEqual(devanagari_font("NirmalaUI"), "Nirmala UI" if os.name == "nt" else "Noto Sans Devanagari")

    def test_ppt_has_one_editable_shape_per_cluster_and_complex_font(self):
        script = Path(__file__).with_name("pdf-to-ppt-build.py")
        spec = importlib.util.spec_from_file_location("ppt_builder", script)
        module = importlib.util.module_from_spec(spec)
        spec.loader.exec_module(module)
        normalized = normalize_devanagari_blocks(blocks(line("कुल राश"), line("ि", (35, 20, 40, 35))))
        presentation = Presentation()
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        module.add_editable_lines(slide, normalized[0]["lines"], 1, 0, 0)
        self.assertEqual(len(slide.shapes), 1)
        self.assertEqual(slide.shapes[0].text, "कुल राशि")
        self.assertIn('lang="hi-IN"', slide.shapes[0]._element.xml)
        self.assertIn("<a:cs", slide.shapes[0]._element.xml)

    def test_word_uses_editable_unicode_and_complex_script_size(self):
        from pdf2docx.text.TextSpan import TextSpan
        install_pdf2docx_devanagari_support()
        document = Document()
        paragraph = document.add_paragraph()
        text_span = TextSpan(span("कुल राशि"))
        text_span.make_docx(paragraph)
        self.assertEqual(paragraph.text, "कुल राशि")
        xml = paragraph._element.xml
        self.assertIn('w:bidi="hi-IN"', xml)
        self.assertIn('w:szCs w:val="24"', xml)
        self.assertIn('w:cs="', xml)


class GridTests(unittest.TestCase):
    def make_table(self):
        table = Document().add_table(rows=2, cols=2)
        table.autofit = False
        for row in table.rows:
            row.cells[0].width = Inches(4)
            row.cells[1].width = Inches(1)
        return table

    def test_consistent_source_widths_replace_equal_default_grid(self):
        table = self.make_table()
        self.assertTrue(synchronize_fixed_table_grid(table))
        self.assertEqual([column.width for column in table.columns], [Inches(4), Inches(1)])
        self.assertEqual(table._tbl.tblPr.find(qn("w:tblW")).get(qn("w:w")), "7200")

    def test_inconsistent_widths_are_not_guessed(self):
        table = self.make_table()
        table.cell(1, 0).width = Inches(3)
        before = table._tbl.xml
        self.assertFalse(synchronize_fixed_table_grid(table))
        self.assertEqual(table._tbl.xml, before)

    def test_merged_table_is_not_changed(self):
        table = self.make_table()
        table.cell(0, 0).merge(table.cell(0, 1))
        before = table._tbl.xml
        self.assertFalse(synchronize_fixed_table_grid(table))
        self.assertEqual(table._tbl.xml, before)


if __name__ == "__main__":
    unittest.main()
