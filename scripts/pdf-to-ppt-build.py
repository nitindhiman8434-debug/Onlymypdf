"""Preserve PDF visuals and place selectable source text on editable slides.

Text is removed only from the rendered background and recreated as on-slide
PowerPoint text boxes. Image-only scans remain visual slides, not fake text.
"""
from __future__ import annotations

import json
import os
import re
import sys
from concurrent.futures import ProcessPoolExecutor, as_completed
from multiprocessing import cpu_count

try:
    import fitz
except ImportError:
    print(json.dumps({"error": "pymupdf not installed. Run: pip install pymupdf"}))
    sys.exit(1)

try:
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.enum.text import MSO_AUTO_SIZE, MSO_ANCHOR
    from pptx.dml.color import RGBColor
    from PIL import Image
except ImportError:
    print(json.dumps({"error": "python-pptx and Pillow required. Run: pip install python-pptx pillow"}))
    sys.exit(1)

PT_PER_INCH = 72.0
MAX_EDITABLE_LINES_PER_PAGE = 1500
MAX_RELIABLE_ON_SLIDE_LINES = 70


def resolve_target_width(page_count: int) -> int:
    if page_count <= 30:
        return 1920
    if page_count <= 80:
        return 1440
    if page_count <= 200:
        return 1280
    if page_count <= 400:
        return 960
    return 720


def resolve_image_format(page_count: int) -> str:
    return "JPEG" if page_count > 80 else "PNG"


def resolve_jpeg_quality(page_count: int) -> int:
    """Lower quality slightly for long documents — visually identical on slides, much smaller files."""
    if page_count <= 80:
        return 85
    if page_count <= 200:
        return 80
    if page_count <= 400:
        return 76
    return 72


def save_slide_image(pix, out_path: str, image_format: str, jpeg_quality: int) -> None:
    if image_format == "JPEG":
        img = Image.frombytes("RGB", (pix.width, pix.height), pix.samples)
        img.save(
            out_path,
            format="JPEG",
            quality=jpeg_quality,
            optimize=True,
            progressive=True,
            subsampling=2,
        )
    else:
        pix.pil_save(out_path, format="PNG", optimize=True)


def pt_to_inches(pt: float) -> float:
    return pt / PT_PER_INCH


def clean_text(value: str) -> str:
    return re.sub(r"[\x00-\x08\x0b\x0c\x0e-\x1f]", "", value)


def resolve_slide_size_inches(sizes: list[tuple[float, float]]) -> tuple[float, float]:
    if not sizes:
        return 10.0, 7.5
    if len(sizes) == 1:
        return sizes[0]
    portrait = sum(1 for w, h in sizes if h >= w)
    landscape = len(sizes) - portrait
    # A tie uses landscape, the more practical uniform canvas for a deck.
    dominant_portrait = portrait > landscape
    pool = [(w, h) for w, h in sizes if (h >= w if dominant_portrait else w > h)] or sizes
    return max(w for w, _ in pool), max(h for _, h in pool)


def extract_page_lines(page: fitz.Page) -> list[dict]:
    lines: list[dict] = []
    data = page.get_text("dict")
    traces = page.get_texttrace()
    visible_boxes = [
        fitz.Rect(span["bbox"])
        for span in traces
        if span.get("type") in (0, 1) and span.get("opacity", 1) > 0.01
    ]
    text_blocks = sorted(
        [b for b in data.get("blocks", []) if b.get("type") == 0],
        key=lambda b: (round(b["bbox"][1], 1), round(b["bbox"][0], 1)),
    )
    for block in text_blocks:
        for line in block.get("lines", []):
            direction = line.get("dir", (1, 0))
            # Rotated text stays in the rendered background until its geometry
            # can be reproduced reliably as an editable PowerPoint shape.
            if abs(direction[0] - 1) > 0.01 or abs(direction[1]) > 0.01:
                continue
            spans = [
                {
                    "text": clean_text(span.get("text", "")),
                    "size": span.get("size", 11),
                    "font": span.get("font", "Arial"),
                    "color": span.get("color", 0),
                }
                for span in line.get("spans", [])
            ]
            bounds = fitz.Rect(line["bbox"])
            if (
                any(span["text"].strip() for span in spans)
                and (not traces or any(bounds.intersects(box) for box in visible_boxes))
            ):
                lines.append({"bbox": line["bbox"], "spans": spans})
    return lines


def is_full_page_raster(page: fitz.Page) -> bool:
    area = page.rect.width * page.rect.height
    if area <= 0:
        return False
    covers_page = any(
        fitz.Rect(image["bbox"]).intersect(page.rect).get_area() / area >= 0.85
        for image in page.get_image_info()
    )
    if not covers_page:
        return False
    # Searchable scans often have an invisible OCR layer. A slide built from
    # that layer would draw duplicate text over the photographed text. A photo
    # background with genuinely visible PDF text is still editable.
    return not any(
        span.get("type") in (0, 1) and span.get("opacity", 1) > 0.01
        for span in page.get_texttrace()
    )


def remove_editable_text_from_background(page: fitz.Page, lines: list[dict]) -> None:
    for line in lines:
        page.add_redact_annot(fitz.Rect(line["bbox"]), fill=False, cross_out=False)
    if lines:
        # Preserve photographs, chart vectors, rules, and page fills.
        page.apply_redactions(images=0, graphics=0, text=0)


def add_editable_lines(slide, lines: list[dict], scale: float, left: float, top: float) -> int:
    count = 0
    for line in lines:
        x0, y0, x1, y1 = line["bbox"]
        size = max(span["size"] for span in line["spans"])
        box = slide.shapes.add_textbox(
            Pt(left + x0 * scale),
            Pt(top + (y0 - size * 0.14) * scale),
            Pt(max(1, (x1 - x0 + size * 0.35) * scale)),
            Pt(max(1, (y1 - y0 + size * 0.45) * scale)),
        )
        box.name = f"Editable line {count + 1}"
        frame = box.text_frame
        frame.clear()
        frame.word_wrap = False
        frame.auto_size = MSO_AUTO_SIZE.NONE
        frame.vertical_anchor = MSO_ANCHOR.TOP
        frame.margin_left = frame.margin_right = frame.margin_top = frame.margin_bottom = 0
        paragraph = frame.paragraphs[0]
        paragraph.space_before = paragraph.space_after = Pt(0)
        for span in line["spans"]:
            run = paragraph.add_run()
            run.text = span["text"]
            run.font.size = Pt(max(1, span["size"] * scale))
            run.font.name = re.sub(r"^[A-Z]{6}\+", "", span["font"])
            if span["color"] is not None:
                color = span["color"]
                run.font.color.rgb = RGBColor(
                    (color >> 16) & 255, (color >> 8) & 255, color & 255
                )
            run.font.bold = "bold" in span["font"].lower()
            run.font.italic = (
                "italic" in span["font"].lower() or "oblique" in span["font"].lower()
            )
        count += 1
    return count


def render_one_page(
    pdf_path: str,
    page_index: int,
    target_width: int,
    output_dir: str,
    image_format: str,
    jpeg_quality: int,
) -> dict:
    doc = fitz.open(pdf_path)
    try:
        page = doc[page_index]
        lines = [] if is_full_page_raster(page) else extract_page_lines(page)
        if len(lines) > MAX_EDITABLE_LINES_PER_PAGE:
            raise ValueError(f"Page {page_index + 1} has too many editable text lines")
        # Dense pages wrap unpredictably in PowerPoint/LibreOffice fonts. Keep
        # their visual page intact and place the selectable transcript in notes.
        visual_reference = len(lines) > MAX_RELIABLE_ON_SLIDE_LINES
        transcript = page.get_text("text", sort=True).strip() if visual_reference else ""
        if not visual_reference:
            remove_editable_text_from_background(page, lines)
        scale = target_width / page.rect.width
        pix = page.get_pixmap(matrix=fitz.Matrix(scale, scale), alpha=False)
        ext = "jpg" if image_format == "JPEG" else "png"
        out_path = os.path.join(output_dir, f"slide{page_index + 1:04d}.{ext}")
        save_slide_image(pix, out_path, image_format, jpeg_quality)
        return {
            "page": page_index + 1,
            "path": out_path,
            "widthPt": round(page.rect.width, 4),
            "heightPt": round(page.rect.height, 4),
            "lines": lines,
            "visualReference": visual_reference,
            "transcript": transcript,
        }
    finally:
        doc.close()


def build_pptx(pdf_path: str, pages: list[dict], output_path: str) -> tuple[int, int]:
    sizes = [(pt_to_inches(p["widthPt"]), pt_to_inches(p["heightPt"])) for p in pages]
    slide_w_in, slide_h_in = resolve_slide_size_inches(sizes)

    prs = Presentation()
    prs.slide_width = Inches(slide_w_in)
    prs.slide_height = Inches(slide_h_in)
    blank = prs.slide_layouts[6]

    editable_slides = 0
    editable_notes_slides = 0
    for page in pages:
        slide = prs.slides.add_slide(blank)
        scale = min(
            slide_w_in * PT_PER_INCH / page["widthPt"],
            slide_h_in * PT_PER_INCH / page["heightPt"],
        )
        left = (slide_w_in * PT_PER_INCH - page["widthPt"] * scale) / 2
        top = (slide_h_in * PT_PER_INCH - page["heightPt"] * scale) / 2
        slide.shapes.add_picture(
            page["path"],
            Pt(left),
            Pt(top),
            width=Pt(page["widthPt"] * scale),
            height=Pt(page["heightPt"] * scale),
        )
        if page["visualReference"]:
            slide.notes_slide.notes_text_frame.text = page["transcript"]
            editable_notes_slides += 1
        elif add_editable_lines(slide, page["lines"], scale, left, top):
            editable_slides += 1

    prs.save(output_path)
    return editable_slides, editable_notes_slides


def main() -> None:
    if len(sys.argv) < 4:
        print(json.dumps({"error": "Usage: python pdf-to-ppt-build.py <input.pdf> <slides-dir> <output.pptx>"}))
        sys.exit(1)

    pdf_path = sys.argv[1]
    output_dir = sys.argv[2]
    output_path = sys.argv[3]

    if not os.path.isfile(pdf_path):
        print(json.dumps({"error": f"File not found: {pdf_path}"}))
        sys.exit(1)

    os.makedirs(output_dir, exist_ok=True)

    doc = fitz.open(pdf_path)
    page_count = doc.page_count
    doc.close()

    if page_count == 0:
        print(json.dumps({"error": "No pages found in the PDF file."}))
        sys.exit(1)

    target_width = resolve_target_width(page_count)
    image_format = resolve_image_format(page_count)
    jpeg_quality = resolve_jpeg_quality(page_count)
    workers = max(1, min(4, cpu_count() or 1))

    pages: list[dict] = []
    tasks = [
        (pdf_path, i, target_width, output_dir, image_format, jpeg_quality)
        for i in range(page_count)
    ]

    with ProcessPoolExecutor(max_workers=workers) as pool:
        futures = [pool.submit(render_one_page, *task) for task in tasks]
        for future in as_completed(futures):
            pages.append(future.result())

    pages.sort(key=lambda p: p["page"])
    editable_slides, editable_notes_slides = build_pptx(pdf_path, pages, output_path)

    print(
        json.dumps(
            {
                "pageCount": page_count,
                "editableSlides": editable_slides,
                "editableNotesSlides": editable_notes_slides,
                "targetWidth": target_width,
                "imageFormat": image_format,
                "jpegQuality": jpeg_quality,
                "workers": workers,
                "outputPath": output_path,
                "outputSizeBytes": os.path.getsize(output_path),
                "mode": "adaptive_on_slide_or_visual_reference_with_editable_notes",
            }
        )
    )


if __name__ == "__main__":
    main()
