"""Generate and inspect source-visible PDF-to-PowerPoint semantic fixtures."""
from __future__ import annotations

import json
import os
import sys

import fitz
from PIL import Image, ImageChops, ImageStat
from pptx import Presentation


def save(doc, out_dir, case_id):
    path = os.path.join(out_dir, case_id + ".pdf")
    doc.save(path, garbage=4, deflate=True)
    for number, page in enumerate(doc):
        page.get_pixmap(matrix=fitz.Matrix(1.5, 1.5), alpha=False).save(
            os.path.join(out_dir, f"{case_id}-source-{number + 1}.png")
        )
    doc.close()


def make_cases(out_dir):
    os.makedirs(out_dir, exist_ok=True)
    cases = []

    doc = fitz.open()
    page = doc.new_page(width=720, height=405)
    page.draw_rect(page.rect, fill=(0.94, 0.97, 1), color=None)
    page.draw_rect(fitz.Rect(40, 40, 680, 355), color=(0.0, 0.35, 0.7))
    page.insert_text((65, 95), "Quarterly Results", fontsize=30, color=(0.1, 0.2, 0.5))
    page.insert_text((65, 160), "Revenue 2026: 123456", fontsize=18)
    page.insert_text((65, 225), "North | South | East", fontsize=16)
    save(doc, out_dir, "styled-vector")
    cases.append({"id": "styled-vector", "tokens": [["Quarterly Results", "Revenue 2026", "North | South | East"]], "editable": [True]})

    doc = fitz.open()
    for number in range(1, 4):
        page = doc.new_page(width=720, height=405)
        page.draw_rect(fitz.Rect(60, 100, 660, 300), fill=(0.95, 0.95, 0.95))
        page.insert_text((80, 130), f"Report page {number}", fontsize=26)
        page.insert_text((80, 200), f"Unique marker {number * 137}", fontsize=18)
    save(doc, out_dir, "three-pages")
    cases.append({"id": "three-pages", "tokens": [[f"Report page {n}", f"Unique marker {n * 137}"] for n in range(1, 4)], "editable": [True] * 3})

    doc = fitz.open()
    first = doc.new_page(width=612, height=792)
    first.insert_text((70, 120), "Portrait cover 2026", fontsize=22)
    second = doc.new_page(width=792, height=612)
    second.insert_text((80, 100), "Landscape detail 2027", fontsize=22)
    save(doc, out_dir, "mixed-orientation")
    cases.append({"id": "mixed-orientation", "tokens": [["Portrait cover 2026"], ["Landscape detail 2027"]], "editable": [True, True]})

    art = fitz.open()
    page = art.new_page(width=720, height=405)
    page.draw_rect(page.rect, fill=(0.13, 0.29, 0.48), color=None)
    page.draw_rect(fitz.Rect(30, 300, 690, 350), fill=(0.9, 0.5, 0.2))
    background = page.get_pixmap(matrix=fitz.Matrix(1.5, 1.5), alpha=False).tobytes("png")
    art.close()
    doc = fitz.open()
    page = doc.new_page(width=720, height=405)
    page.insert_image(page.rect, stream=background)
    page.insert_text((60, 120), "Text on photograph", fontsize=28, color=(1, 1, 1))
    save(doc, out_dir, "photo-with-text")
    cases.append({"id": "photo-with-text", "tokens": [["Text on photograph"]], "editable": [True]})

    art = fitz.open()
    page = art.new_page(width=720, height=405)
    page.insert_text((70, 140), "Scanned poster 2040", fontsize=32)
    poster = page.get_pixmap(matrix=fitz.Matrix(2, 2), alpha=False).tobytes("png")
    art.close()
    doc = fitz.open()
    page = doc.new_page(width=720, height=405)
    page.insert_image(page.rect, stream=poster)
    save(doc, out_dir, "image-only")
    cases.append({"id": "image-only", "tokens": [[]], "editable": [False]})

    doc = fitz.open()
    page = doc.new_page(width=720, height=405)
    page.insert_image(page.rect, stream=poster)
    page.insert_text((70, 140), "Scanned poster 2040", fontsize=32, render_mode=3)
    save(doc, out_dir, "hidden-ocr")
    cases.append({"id": "hidden-ocr", "tokens": [[]], "editable": [False]})

    doc = fitz.open()
    page = doc.new_page(width=720, height=405)
    page.insert_text((65, 100), "Horizontal editable", fontsize=22)
    page.insert_text((650, 320), "Vertical label", fontsize=18, rotate=90)
    save(doc, out_dir, "rotated-label")
    cases.append({"id": "rotated-label", "tokens": [["Horizontal editable"]], "editable": [True]})

    print(json.dumps({"cases": cases}))


def inspect_case(case, out_dir):
    case_id = case["id"]
    pptx = Presentation(os.path.join(out_dir, case_id + ".pptx"))
    source = fitz.open(os.path.join(out_dir, case_id + ".pdf"))
    rendered = fitz.open(os.path.join(out_dir, "rendered", case_id + ".pdf"))
    assert len(pptx.slides) == len(source) == len(rendered), f"{case_id}: page count"
    slide_width = pptx.slide_width / 12700
    slide_height = pptx.slide_height / 12700
    errors = []
    visual_scores = []
    text_counts = []

    for index, (slide, source_page, render_page) in enumerate(zip(pptx.slides, source, rendered)):
        text_shapes = [
            shape for shape in slide.shapes
            if shape.has_text_frame and shape.text.strip()
            and shape.left >= 0 and shape.top >= 0
            and shape.left + shape.width <= pptx.slide_width + 12700
            and shape.top + shape.height <= pptx.slide_height + 12700
        ]
        texts = " ".join(shape.text for shape in text_shapes)
        text_counts.append(len(text_shapes))
        for token in case["tokens"][index]:
            if token not in texts or token not in render_page.get_text():
                errors.append(f"page {index + 1} lost editable/visible token: {token}")
        if bool(text_shapes) != case["editable"][index]:
            errors.append(f"page {index + 1} editable state mismatch")

        pictures = [shape for shape in slide.shapes if shape.shape_type == 13]
        if len(pictures) != 1:
            errors.append(f"page {index + 1} missing visual background")
        else:
            picture = pictures[0]
            picture_ratio = picture.width / picture.height
            source_ratio = source_page.rect.width / source_page.rect.height
            if abs(picture_ratio - source_ratio) > 0.005:
                errors.append(f"page {index + 1} stretched page image")

        output_pix = render_page.get_pixmap(matrix=fitz.Matrix(1.5, 1.5), alpha=False)
        output_image = Image.frombytes(
            "RGB", (output_pix.width, output_pix.height), output_pix.samples
        )
        expected = Image.new("RGB", output_image.size, "white")
        page_width = source_page.rect.width
        page_height = source_page.rect.height
        scale = min(slide_width / page_width, slide_height / page_height)
        content_width = round(page_width * scale * 1.5)
        content_height = round(page_height * scale * 1.5)
        source_pix = source_page.get_pixmap(matrix=fitz.Matrix(1.5, 1.5), alpha=False)
        original = Image.frombytes("RGB", (source_pix.width, source_pix.height), source_pix.samples)
        original = original.resize((content_width, content_height), Image.Resampling.LANCZOS)
        expected.paste(
            original,
            ((output_image.width - content_width) // 2, (output_image.height - content_height) // 2),
        )
        score = sum(ImageStat.Stat(ImageChops.difference(expected, output_image)).mean) / 3
        visual_scores.append(round(score, 2))
        if score > 15:
            errors.append(f"page {index + 1} visual difference {score:.2f} > 15")
        output_image.save(os.path.join(out_dir, "rendered", f"{case_id}-page-{index + 1}.png"))

    source.close()
    rendered.close()
    return {
        "id": case_id,
        "passed": not errors,
        "pageCount": len(pptx.slides),
        "editableTextBoxes": text_counts,
        "visualMeanAbs": visual_scores,
        "errors": errors,
    }


if __name__ == "__main__":
    if len(sys.argv) < 3:
        raise SystemExit("Usage: phase2-ppt-corpus.py generate|verify <output-dir> [metadata.json]")
    if sys.argv[1] == "generate":
        make_cases(sys.argv[2])
    elif sys.argv[1] == "verify":
        with open(sys.argv[3], encoding="utf-8") as stream:
            cases = json.load(stream)["cases"]
        print(json.dumps({"results": [inspect_case(case, sys.argv[2]) for case in cases]}))
    else:
        raise SystemExit("Unknown command")
