#!/usr/bin/env python3
"""Bounded English scanned-invoice native Word-table gate; loopback HTTP only."""
from __future__ import annotations

import argparse
import copy
import hashlib
import importlib.metadata
import importlib.util
import io
import json
import platform
import re
import time
import zipfile
from collections import Counter
from datetime import datetime, timezone
from pathlib import Path
from xml.etree import ElementTree as ET

spec = importlib.util.spec_from_file_location("english_table_word_helpers", Path(__file__).with_name("phase3-english-word-corpus.py"))
word = importlib.util.module_from_spec(spec)
spec.loader.exec_module(word)
ROOT, W = word.ROOT, word.W
ARTIFACT_ROOT = ROOT / "tmp" / "pdfs" / "phase3.2f-english-table"
REPORT_ROOT = ROOT / "quality" / "phase3-english-table"
LIMITATIONS = [
    "Four fixed synthetic English-only cases; not a universal table or OCR quality claim.",
    "Three positive cases use complete rectangular black ruling, four unmerged columns and a header plus three body rows per invoice.",
    "The negative memo contains aligned borderless prose; conservative output must preserve it as paragraphs with zero native tables.",
    "Complete cell matrices, leading-zero identifiers, decimal amounts and surrounding text are exact after NFC and existing-whitespace collapse only.",
    "Source DOCX and independently rendered source PDF must pass before any image-only upload is made.",
    "Returned DOCX is independently rendered with LibreOffice; PNG paths are evidence for separate manual visual QA, not an automatic visual pass.",
    "Reports and artifact directories use exclusive run names; baseline and final evidence cannot overwrite one another.",
    "No paid services, downloads, public deployment, non-loopback requests or other-language fixture generation are performed.",
]


def definitions() -> list[dict]:
    def invoice(title, reference, rows, total, closing):
        return [word.p(title, True), word.p("Invoice number " + reference),
                word.p("Customer: Example Research Company"),
                word.p("The following completed services are synthetic test information."),
                {"kind": "table", "rows": [["Item ID", "Description", "Qty", "Line total"], *rows]},
                word.p("Amount due: " + total + " USD"), word.p(closing)]

    simple = invoice("English service invoice", "004218", [
        ["00421", "Analysis package", "2", "240.50"],
        ["00632", "Review service", "3", "135.30"],
        ["00874", "Archive support", "1", "19.95"],
    ], "395.75", "Please keep every service with its item identifier and amount.")
    wrapped = invoice("Detailed service invoice", "005329", [
        ["00542", "Monthly analysis and reporting package for the research team", "2", "280.50"],
        ["00763", "Final review service and document preparation for the project", "3", "165.30"],
        ["00985", "Archive support", "1", "29.95"],
    ], "475.75", "The descriptions continue inside their own cells and retain their amounts.")
    first = invoice("First page service invoice", "006431", [
        ["00124", "Planning package", "2", "220.50"],
        ["00345", "Design review", "3", "125.30"],
        ["00567", "Record storage", "1", "39.95"],
    ], "385.75", "The first invoice ends here before the second page begins.")
    second = invoice("Second page service invoice", "007542", [
        ["00235", "Delivery package", "2", "320.50"],
        ["00456", "Quality review", "3", "145.30"],
        ["00678", "Final support", "1", "49.95"],
    ], "515.75", "This is the final paragraph of the second invoice and the document.")
    memo = [word.p("Reference review memo", True),
            word.p("The following notes are plain paragraphs without a ruled table."),
            word.p("Reference 006731\tApproved amount 81.25 USD"),
            word.p("Reference 008952\tApproved amount 42.60 USD"),
            word.p("Reference 003174\tApproved amount 19.95 USD"),
            word.p("Keep these notes in order and retain the complete final paragraph.")]
    cases = [
        {"caseId": "english-ruled-invoice-scan", "pages": [simple]},
        {"caseId": "english-wrapped-invoice-scan", "pages": [wrapped],
         "requiredWrappedDescriptions": [wrapped[4]["rows"][1][1], wrapped[4]["rows"][2][1]]},
        {"caseId": "english-two-invoice-scan", "pages": [first, second], "pageSizesPoints": [[612, 792], [792, 612]]},
        {"caseId": "english-borderless-memo-scan", "pages": [memo], "negativeControl": True},
    ]
    for case in cases:
        case["inputKind"] = "scan"
        case["numericValues"] = re.findall(r"\d+(?:\.\d+)?", "\n".join(text for page in case["pages"] for text in word.page_text(page)))
    return cases


def expected_tables(case: dict) -> list[list[list[str]]]:
    return [[[word.normalize(cell) for cell in row] for row in block["rows"]]
            for page in case["pages"] for block in page if block["kind"] == "table"]


def inspect_docx(body: bytes, case: dict) -> dict:
    result = word.inspect_docx(body, case)
    with zipfile.ZipFile(io.BytesIO(body)) as archive:
        root = ET.fromstring(archive.read("word/document.xml")).find(W + "body")
    tables = list(root.iter(W + "tbl"))
    matrices = [[[word.normalize(" ".join(word.visible_paragraph(p) for p in cell.iter(W + "p")))
                  for cell in row.findall(W + "tc")] for row in table.findall(W + "tr")] for table in tables]
    expected = expected_tables(case)
    merged = [node.tag.rsplit("}", 1)[-1] for table in tables for node in table.iter()
              if node.tag in (W + "vMerge", W + "hMerge") or (node.tag == W + "gridSpan" and node.get(W + "val", "1") != "1")]
    nested = len(tables) != len(root.findall(W + "tbl"))
    checks = {"passed": matrices == expected and not merged and not nested,
              "expectedTableCount": len(expected), "actualTableCount": len(tables),
              "expectedMatrices": expected, "actualMatrices": matrices,
              "completeMatricesMatch": matrices == expected, "unexpectedMergeProperties": merged,
              "nestedOrWrappedTables": nested, "nativeTableReconstructionRequired": bool(expected),
              "zeroTablesRequired": not expected}
    result["nativeTables"] = checks
    result["passed"] = result["passed"] and checks["passed"]
    return result


def configure_artifacts(directory: Path) -> None:
    directory.resolve().relative_to(ARTIFACT_ROOT.resolve())
    word.ARTIFACTS = directory


def grid_evidence(path: Path, case: dict) -> dict:
    import fitz
    pages, wraps = [], []
    with fitz.open(path) as pdf:
        for index, page in enumerate(pdf):
            horizontal, vertical = [], []
            for drawing in page.get_drawings():
                if drawing.get("color") != (0.0, 0.0, 0.0) or drawing.get("width", 0) < 0.7:
                    continue
                for item in drawing["items"]:
                    if item[0] == "l":
                        a, b = item[1:]
                        if abs(a.y - b.y) < 0.1 and abs(a.x - b.x) > 25:
                            horizontal.append(round(a.y, 1))
                        if abs(a.x - b.x) < 0.1 and abs(a.y - b.y) > 25:
                            vertical.append(round(a.x, 1))
            tables_on_page = sum(block["kind"] == "table" for block in case["pages"][index])
            # Explicit cell borders may be emitted as contiguous segments. Distinct
            # axes establish genuine visible ruling independently of source styles.
            ruled = len(set(horizontal)) >= 5 and len(set(vertical)) >= 5
            pages.append({"page": index + 1, "horizontalAxes": sorted(set(horizontal)), "verticalAxes": sorted(set(vertical)),
                          "expectedRuledTables": tables_on_page, "passed": ruled if tables_on_page else not horizontal and not vertical})
            for phrase in case.get("requiredWrappedDescriptions", []):
                rectangles = page.search_for(phrase)
                if rectangles:
                    line_count = len({round(rect.y0, 1) for rect in rectangles})
                    wraps.append({"text": phrase, "page": index + 1, "lineCount": line_count, "passed": line_count >= 2})
    return {"passed": all(page["passed"] for page in pages) and len(wraps) == len(case.get("requiredWrappedDescriptions", [])) and all(row["passed"] for row in wraps),
            "pages": pages, "wrappedDescriptions": wraps}


def generate_source(case: dict, office: str) -> dict:
    import fitz
    from docx import Document
    from docx.enum.section import WD_SECTION_START
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Pt

    directory = word.ARTIFACTS
    directory.mkdir(parents=True, exist_ok=True)
    source = directory / (case["caseId"] + "-original.docx")
    document = Document()
    font = word.latin_font()
    normal = document.styles["Normal"]
    normal.font.name, normal.font.size = font, Pt(12)
    normal.paragraph_format.space_after, normal.paragraph_format.line_spacing = Pt(10), 1.1
    for page_index, page in enumerate(case["pages"]):
        if page_index:
            document.add_section(WD_SECTION_START.NEW_PAGE)
        section = document.sections[-1]
        width, height = case.get("pageSizesPoints", [[612, 792]] * len(case["pages"]))[page_index]
        section.page_width, section.page_height = Pt(width), Pt(height)
        section.left_margin = section.right_margin = Pt(50.4)
        section.top_margin = section.bottom_margin = Pt(50.4)
        for block in page:
            if block["kind"] != "table":
                paragraph = document.add_paragraph()
                paragraph.paragraph_format.tab_stops.add_tab_stop(Pt(230))
                run = paragraph.add_run(block["text"])
                if block["kind"] == "heading":
                    run.bold, run.font.size = True, Pt(17)
                continue
            table = document.add_table(rows=0, cols=4)
            table.autofit = False
            available = width - 100.8
            widths = [75, available - 225, 50, 100]
            for column, column_width in zip(table.columns, widths):
                column.width = Pt(column_width)
            for row_index, values in enumerate(block["rows"]):
                row = table.add_row()
                for cell, value, cell_width in zip(row.cells, values, widths):
                    cell.width, cell.text = Pt(cell_width), value
                    properties = cell._tc.get_or_add_tcPr()
                    borders = OxmlElement("w:tcBorders")
                    for edge in ("top", "left", "bottom", "right"):
                        border = OxmlElement("w:" + edge)
                        for key, attr in (("val", "single"), ("sz", "8"), ("color", "000000")):
                            border.set(qn("w:" + key), attr)
                        borders.append(border)
                    properties.append(borders)
                    margins = OxmlElement("w:tcMar")
                    for edge in ("top", "left", "bottom", "right"):
                        margin = OxmlElement("w:" + edge)
                        margin.set(qn("w:w"), "110")
                        margin.set(qn("w:type"), "dxa")
                        margins.append(margin)
                    properties.append(margins)
                    for paragraph in cell.paragraphs:
                        paragraph.paragraph_format.space_after = Pt(0)
                        paragraph.paragraph_format.line_spacing = 1.15
                        for run in paragraph.runs:
                            run.bold = row_index == 0
            document.add_paragraph()
    document.save(source)
    original = inspect_docx(source.read_bytes(), case)
    if not original["passed"]:
        raise ValueError("Literal source DOCX exact text/native-table validation failed: " + case["caseId"])
    original_pdf = word.render_docx(source, directory / "original-pdfs", office, "source-libreoffice-profile")
    source_render = word.inspect_render(original_pdf, case, case["caseId"] + "-source-preview")
    ruling = grid_evidence(original_pdf, case)
    if not source_render["passed"] or not ruling["passed"]:
        raise ValueError("Original PDF text, pagination, visible ruling or required wrapping failed: " + json.dumps({"caseId": case["caseId"], "render": source_render, "ruling": ruling}))
    upload = directory / (case["caseId"] + ".pdf")
    with fitz.open(original_pdf) as original_pdf_doc, fitz.open() as scan:
        for index, page in enumerate(original_pdf_doc, 1):
            raster = page.get_pixmap(dpi=300, colorspace=fitz.csRGB, alpha=False)
            png = directory / f"{case['caseId']}-upload-image-{index}.png"
            png.write_bytes(raster.tobytes("png"))
            scan.new_page(width=page.rect.width, height=page.rect.height).insert_image(page.rect, filename=str(png))
        scan.save(upload, garbage=4, deflate=True)
    with fitz.open(upload) as pdf:
        counts = [len(page.get_text().strip()) for page in pdf]
        images = [len(page.get_images()) for page in pdf]
        valid = len(pdf) == len(case["pages"]) and all(count == 0 for count in counts) and all(count == 1 for count in images)
    return {"passed": valid, "inputPath": word.relative(upload), "inputSha256": word.digest(upload), "inputBytes": upload.stat().st_size,
            "inputKind": "scan", "selectableCharactersByPage": counts, "imageCountByPage": images, "fontFamily": font,
            "literalDocxPath": word.relative(source), "literalDocxSha256": word.digest(source),
            "literalDocxChecks": original, "trustedOriginal": source_render, "sourceRulingAndWrapping": ruling}


def convert(case: dict, base: str, timeout: float, office: str) -> dict:
    start = time.monotonic()
    result = {"caseId": case["caseId"], "tool": "pdf-to-word"}
    try:
        transport = word.transport_module.LocalTransport(base, timeout)
        payload, content_type = word.multipart(ROOT / case["source"]["inputPath"])
        status, body, headers = transport.request("/api/tools/pdf-to-word", payload, content_type)
        mime = headers.get("content-type", "").split(";", 1)[0].strip().lower()
        result.update({"httpStatus": status, "mime": mime, "outputBytes": len(body), "responseSha256": hashlib.sha256(body).hexdigest(),
                       "responseEngine": headers.get("x-pdf-engine"), "conversionDurationMs": round((time.monotonic() - start) * 1000)})
        if status != 200 or mime != word.DOCX_MIME:
            result.update({"status": "failed", "reason": "Expected HTTP 200 with DOCX MIME"})
            return result
        output = word.ARTIFACTS / (case["caseId"] + "-result.docx")
        output.write_bytes(body)
        result["outputPath"] = word.relative(output)
        result["document"] = inspect_docx(body, case)
        rendered = word.render_docx(output, word.ARTIFACTS / "rendered-outputs", office, "output-libreoffice-profile")
        result["rendered"] = word.inspect_render(rendered, case, case["caseId"] + "-output-preview")
        result["status"] = "passed" if result["document"]["passed"] and result["rendered"]["passed"] else "failed"
    except Exception as exc:
        result.update({"status": "error", "error": str(exc)[:2000]})
    finally:
        result["totalDurationMs"] = round((time.monotonic() - start) * 1000)
    return result


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--base-url", default="http://127.0.0.1:3001")
    parser.add_argument("--generate-only", action="store_true")
    parser.add_argument("--run-label", required=True, help="Fresh immutable run identifier, e.g. baseline or final")
    parser.add_argument("--case", action="append", choices=[case["caseId"] for case in definitions()])
    parser.add_argument("--report", type=Path)
    parser.add_argument("--reuse-sources", type=Path, help="Existing generate-only report; reuse its validated, hash-checked scans")
    parser.add_argument("--timeout", type=float, default=180)
    parser.add_argument("--libreoffice")
    args = parser.parse_args()
    if not re.fullmatch(r"[a-z0-9][a-z0-9-]{0,63}", args.run_label):
        parser.error("run-label must be 1-64 lowercase letters, numbers or hyphens")
    if not 1 <= args.timeout <= 600:
        parser.error("timeout must be between 1 and 600 seconds")
    if args.generate_only and args.reuse_sources:
        parser.error("generate-only and reuse-sources cannot be combined")
    base = word.transport_module.validate_base_url(args.base_url)
    directory = ARTIFACT_ROOT / args.run_label
    report_path = args.report or REPORT_ROOT / (args.run_label + "-report.json")
    if report_path.exists() or directory.exists():
        parser.error("Run/report already exists; use a fresh run-label/report to preserve immutable evidence")
    selected = [case for case in definitions() if not args.case or case["caseId"] in args.case]
    report = {"schemaVersion": 1, "phase": "3.2F", "runLabel": args.run_label, "generatedAt": datetime.now(timezone.utc).isoformat(),
              "baseUrl": base, "generateOnly": args.generate_only, "limitations": LIMITATIONS,
              "validatorSha256": word.digest(Path(__file__)), "fixtureDefinitionsSha256": hashlib.sha256(json.dumps(definitions(), sort_keys=True).encode()).hexdigest(),
              "manualVisualQa": {"status": "required", "instructions": "Inspect every source/output PNG for borders, cell association, clipping, wrapping and page completeness."},
              "environment": {"python": platform.python_version(), "platform": platform.system(),
                              "packages": {name: importlib.metadata.version(name) for name in ("PyMuPDF", "python-docx", "Pillow")}},
              "cases": [], "results": []}
    configure_artifacts(directory)
    directory.mkdir(parents=True, exist_ok=False)
    try:
        office = word.find_office(args.libreoffice)
        sources = None
        if args.reuse_sources:
            previous = json.loads(args.reuse_sources.read_text(encoding="utf-8"))
            if not previous.get("generateOnly") or previous.get("fixtureDefinitionsSha256") != report["fixtureDefinitionsSha256"]:
                raise ValueError("Reused sources require a generate-only report with identical literal definitions")
            sources = {item["caseId"]: item for item in previous["cases"]}
            report["reusedSourceReport"] = {"path": str(args.reuse_sources), "sha256": word.digest(args.reuse_sources)}
        for case in selected:
            if sources is None:
                source = generate_source(case, office)
            else:
                source = copy.deepcopy(sources[case["caseId"]]["source"])
                if not source["passed"]:
                    raise ValueError("Cannot reuse invalid source")
                for path_key, hash_key in (("inputPath", "inputSha256"), ("literalDocxPath", "literalDocxSha256")):
                    if word.digest(ROOT / source[path_key]) != source[hash_key]:
                        raise ValueError("Source artifact hash changed: " + path_key)
                original = source["trustedOriginal"]
                if word.digest(ROOT / original["pdfPath"]) != original["pdfSha256"]:
                    raise ValueError("Trusted original PDF hash changed")
            item = {**case, "source": source}
            report["cases"].append(item)
            if source["passed"] and not args.generate_only:
                result = convert(item, base, args.timeout, office)
                report["results"].append(result)
                print(json.dumps({"caseId": case["caseId"], "status": result["status"], "httpStatus": result.get("httpStatus")}), flush=True)
    except Exception as exc:
        report["setupError"] = str(exc)[:12000]
    failed_sources = sum(not case["source"]["passed"] for case in report["cases"])
    statuses = Counter(result["status"] for result in report["results"])
    complete = len(report["cases"]) == len(selected) and not failed_sources and not report.get("setupError")
    passed = complete and len(report["results"]) == len(selected) and statuses["passed"] == len(selected)
    report["summary"] = {"casesSelected": len(selected), "fixturesValidated": len(report["cases"]), "fixturesFailed": failed_sources,
                         "routesTested": len(report["results"]), "passed": statuses["passed"], "failed": statuses["failed"], "errors": statuses["error"],
                         "allPassed": None if args.generate_only else bool(passed)}
    report_path.parent.mkdir(parents=True, exist_ok=True)
    with report_path.open("x", encoding="utf-8") as handle:
        handle.write(json.dumps(report, ensure_ascii=False, indent=2) + "\n")
    print(json.dumps({"report": str(report_path), "summary": report["summary"]}), flush=True)
    return 0 if (complete if args.generate_only else passed) else 1


if __name__ == "__main__":
    raise SystemExit(main())
