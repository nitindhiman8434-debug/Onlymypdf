#!/usr/bin/env python3
"""Local HTTP Office-conversion QA against separately downloaded public PDFs.

Source PDFs and output files stay in the ignored real-world/ directory. The
checked-in JSON contains only measurements, not document contents.
"""

from __future__ import annotations

import hashlib
import http.cookiejar
import json
import os
import re
import time
import urllib.error
import urllib.request
import uuid
import zipfile
from datetime import datetime, timezone
from pathlib import Path
from xml.etree import ElementTree

import fitz
from openpyxl import load_workbook
from pptx import Presentation


ROOT = Path(__file__).resolve().parents[1]
CORPUS = ROOT / "quality" / "phase2-production" / "real-world"
REPORT = ROOT / "quality" / "phase2-production" / "local-real-world-report.json"
BASE_URL = "http://127.0.0.1:3001"
SAMPLES = [
    ("irs-form-1040-2025.pdf", "form"),
    ("arxiv-2010.12647.pdf", "two-column-paper"),
    ("archives-declaration-scan.pdf", "image-only-scan"),
]
TOOLS = [
    ("pdf-to-word", ".docx"),
    ("pdf-to-excel", ".xlsx"),
    ("pdf-to-ppt", ".pptx"),
]
XML_WORD_TEXT = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}t"


def utc_now() -> str:
    return datetime.now(timezone.utc).isoformat()


def tokens(value: str) -> set[str]:
    return set(re.findall(r"[a-z0-9]{4,}", value.lower()))


def multipart(pdf_bytes: bytes, name: str, word: bool) -> tuple[bytes, str]:
    boundary = f"onlymypdf-qa-{uuid.uuid4().hex}"
    body = (
        f"--{boundary}\r\n"
        f'Content-Disposition: form-data; name="file"; filename="{name}"\r\n'
        "Content-Type: application/pdf\r\n\r\n"
    ).encode() + pdf_bytes + f"\r\n--{boundary}".encode()
    if word:
        body += b'\r\nContent-Disposition: form-data; name="options"\r\n\r\n{}'
        body += f"\r\n--{boundary}".encode()
    body += b"--\r\n"
    return body, f"multipart/form-data; boundary={boundary}"


def request(opener, url: str, *, body: bytes | None = None, content_type: str | None = None,
            word: bool = False) -> tuple[int, bytes, dict[str, str]]:
    headers = {"Origin": BASE_URL, "User-Agent": "OnlyMyPDF-local-QA/1"}
    if content_type:
        headers["Content-Type"] = content_type
    if word:
        headers["X-Pdf-To-Word-Job"] = "1"
    req = urllib.request.Request(url, data=body, headers=headers, method="POST" if body is not None else "GET")
    try:
        with opener.open(req, timeout=300) as response:
            return response.status, response.read(), {key.lower(): value for key, value in response.headers.items()}
    except urllib.error.HTTPError as error:
        return error.code, error.read(), {key.lower(): value for key, value in error.headers.items()}


def convert(pdf_path: Path, tool: str, extension: str) -> tuple[int, bytes, dict[str, str], str | None, float]:
    opener = urllib.request.build_opener(urllib.request.HTTPCookieProcessor(http.cookiejar.CookieJar()))
    payload, content_type = multipart(pdf_path.read_bytes(), pdf_path.name, tool == "pdf-to-word")
    started = time.monotonic()
    status, data, headers = request(opener, f"{BASE_URL}/api/tools/{tool}", body=payload,
                                    content_type=content_type, word=tool == "pdf-to-word")
    engine = None
    if tool == "pdf-to-word" and status == 202:
        job_id = json.loads(data)["jobId"]
        for _ in range(240):
            poll_status, poll_data, _ = request(
                opener, f"{BASE_URL}/api/tools/pdf-to-word/status?jobId={job_id}"
            )
            if poll_status != 200:
                return poll_status, poll_data, {}, engine, time.monotonic() - started
            state = json.loads(poll_data)
            if state.get("status") == "error":
                return 500, json.dumps(state).encode(), {}, engine, time.monotonic() - started
            if state.get("status") == "done":
                engine = state.get("engine")
                status, data, headers = request(
                    opener, f"{BASE_URL}/api/tools/pdf-to-word/download?jobId={job_id}"
                )
                break
            time.sleep(1)
        else:
            return 504, b"Word job polling timed out", {}, engine, time.monotonic() - started
    return status, data, headers, engine, time.monotonic() - started


def source_metrics(pdf_path: Path) -> dict:
    with fitz.open(pdf_path) as pdf:
        text = "\n".join(page.get_text() for page in pdf)
        return {
            "pages": len(pdf),
            "bytes": pdf_path.stat().st_size,
            "sha256": hashlib.sha256(pdf_path.read_bytes()).hexdigest(),
            "selectableTextChars": len(text),
            "imageCount": sum(len(page.get_images()) for page in pdf),
            "tokenSet": tokens(text),
        }


def inspect_office(output_path: Path, extension: str, source_tokens: set[str]) -> dict:
    with zipfile.ZipFile(output_path) as archive:
        bad_part = archive.testzip()
        if bad_part or "[Content_Types].xml" not in archive.namelist():
            raise ValueError(f"Invalid Office ZIP package: {bad_part or 'content types missing'}")
        if extension == ".docx":
            root = ElementTree.fromstring(archive.read("word/document.xml"))
            text = " ".join(node.text or "" for node in root.iter(XML_WORD_TEXT))
            structure = {"mediaParts": len([n for n in archive.namelist() if n.startswith("word/media/")])}
        elif extension == ".xlsx":
            workbook = load_workbook(output_path, read_only=True, data_only=False)
            cells = [str(cell.value) for sheet in workbook for row in sheet for cell in row if cell.value is not None]
            text = " ".join(cells)
            structure = {"sheets": len(workbook.sheetnames), "populatedCells": len(cells)}
            workbook.close()
        else:
            presentation = Presentation(output_path)
            slide_text = [shape.text for slide in presentation.slides for shape in slide.shapes
                          if shape.has_text_frame]
            notes_text = [slide.notes_slide.notes_text_frame.text for slide in presentation.slides
                          if slide.has_notes_slide and slide.notes_slide.notes_text_frame]
            text = " ".join(slide_text + notes_text)
            structure = {
                "slides": len(presentation.slides),
                "pictures": sum(shape.shape_type == 13 for slide in presentation.slides for shape in slide.shapes),
                "editableTextShapes": sum(bool(shape.text.strip()) for slide in presentation.slides
                                          for shape in slide.shapes if shape.has_text_frame),
                "onSlideTextChars": len(" ".join(slide_text)),
                "notesTextChars": len(" ".join(notes_text)),
            }
    output_tokens = tokens(text)
    return {
        "zipCrcPassed": True,
        "editableTextChars": len(text),
        "uniqueTokenRecall": round(len(source_tokens & output_tokens) / len(source_tokens), 4)
        if source_tokens else None,
        **structure,
    }


def main() -> None:
    CORPUS.mkdir(parents=True, exist_ok=True)
    selected_sample = os.environ.get("PHASE2_QA_SAMPLE")
    selected_tool = os.environ.get("PHASE2_QA_TOOL")
    report_path = Path(os.environ.get("PHASE2_QA_REPORT", str(REPORT)))
    report = {"startedAt": utc_now(), "baseUrl": BASE_URL,
              "note": "Public samples, local dev HTTP only. Token recall is a diagnostic, not layout or universal accuracy.",
              "cases": []}
    for name, sample_type in SAMPLES:
        if selected_sample and selected_sample != name:
            continue
        pdf_path = CORPUS / name
        if not pdf_path.is_file():
            raise FileNotFoundError(f"Download the public sample first: {pdf_path}")
        source = source_metrics(pdf_path)
        case = {"source": name, "type": sample_type,
                "pages": source["pages"], "inputBytes": source["bytes"],
                "sha256": source["sha256"], "selectableTextChars": source["selectableTextChars"],
                "imageCount": source["imageCount"], "results": []}
        report["cases"].append(case)
        for tool, extension in TOOLS:
            if selected_tool and selected_tool != tool:
                continue
            result = {"tool": tool}
            try:
                status, body, headers, engine, elapsed = convert(pdf_path, tool, extension)
                result.update({"httpStatus": status, "elapsedSeconds": round(elapsed, 2), "engine": engine})
                if status != 200:
                    result["error"] = body.decode("utf-8", "replace")[:300]
                else:
                    output_path = CORPUS / f"{pdf_path.stem}-{tool}{extension}"
                    output_path.write_bytes(body)
                    result.update({"outputBytes": len(body),
                                   "contentType": headers.get("content-type"),
                                   "sha256": hashlib.sha256(body).hexdigest(),
                                   "validation": inspect_office(output_path, extension, source["tokenSet"])})
            except Exception as error:
                result["error"] = f"{type(error).__name__}: {error}"
            case["results"].append(result)
            print(f"{name} {tool}: HTTP {result.get('httpStatus', '?')} {result.get('elapsedSeconds', '?')}s", flush=True)
    report["completedAt"] = utc_now()
    report_path.write_text(json.dumps(report, indent=2) + "\n", encoding="utf-8")
    print(f"Saved {report_path}")


if __name__ == "__main__":
    main()
