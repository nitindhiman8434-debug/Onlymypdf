"""Word-cell and same-pass OCR contracts independent of a real recognizer."""

from pathlib import Path
import copy
import tempfile
from types import SimpleNamespace
import unittest
from unittest.mock import patch

import pymupdf as fitz
from docx import Document

from ocr_english_docx import recognize_ocr_transcripts, write_english_ocr_docx


class TableIntegrationTests(unittest.TestCase):
    def setUp(self):
        directory = tempfile.TemporaryDirectory()
        self.addCleanup(directory.cleanup)
        self.source = Path(directory.name) / "source.pdf"
        self.output = Path(directory.name) / "result.docx"
        with fitz.open() as document:
            document.new_page(width=612, height=792)
            document.new_page(width=792, height=612)
            document.save(self.source)
        self.layout = {"before": ["Invoice 001234"],
                       "rows": [["Description", "Amount"], ["Repeated service", "120.25"],
                                ["Repeated service", "120.25"]],
                       "after": ["Total 240.50"], "columnWidths": [0.75, 0.25]}
        self.raw = "Invoice 001234 Description Amount Repeated service 120.25 Repeated service 120.25 Total 240.50"

    def test_native_cells_preserve_literals_repeats_body_order_and_page_geometry(self):
        write_english_ocr_docx(str(self.source), str(self.output), [self.raw, "Last page reference"],
                              page_layouts=[self.layout, None])
        document = Document(self.output)
        self.assertEqual(len(document.tables), 1)
        table = document.tables[0]
        self.assertEqual([[c.text for c in row.cells] for row in table.rows], self.layout["rows"])
        self.assertEqual(len(document.inline_shapes), 0)
        self.assertEqual([p.text for p in document.paragraphs if p.text],
                         ["Invoice 001234", "Total 240.50", "Last page reference"])
        body = document.element.body
        blocks = [(el.tag.rsplit("}", 1)[-1], " ".join(el.xpath(".//w:t/text()"))) for el in body]
        nonempty = [(tag, text) for tag, text in blocks if text]
        self.assertEqual([tag for tag, _ in nonempty], ["p", "tbl", "p", "p"])
        self.assertEqual([(s.page_width.pt, s.page_height.pt) for s in document.sections],
                         [(612, 792), (792, 612)])
        self.assertAlmostEqual(table.columns[0].width.pt / table.columns[1].width.pt, 3, places=2)
        self.assertAlmostEqual(sum(c.width.pt for c in table.columns), 612 - 86.4, places=1)

    def test_table_count_mismatch_cannot_silently_drop_later_layout(self):
        with self.assertRaisesRegex(ValueError, "layout count"):
            write_english_ocr_docx(str(self.source), str(self.output), [self.raw, "Last page"],
                                  page_layouts=[self.layout])
        self.assertFalse(self.output.exists())

    def two_table_layout(self):
        return {"blocks": [
            {"kind": "paragraphs", "lines": ["  Statement  001234  ", ""]},
            {"kind": "table", "rows": self.layout["rows"], "columnWidths": [0.75, 0.25]},
            {"kind": "paragraphs", "lines": ["", "  The  next table lists adjustments.  ", ""]},
            {"kind": "table", "rows": [["Code", "Reason", "Amount"], ["0007", "Adjustment", "120.25"]],
             "columnWidths": [0.2, 0.5, 0.3]},
            {"kind": "paragraphs", "lines": ["", "  Statement  001234 ends here.  "]},
        ]}

    @staticmethod
    def layout_text(layout):
        return "\n".join(text for block in layout["blocks"]
                         for text in (block["lines"] if block["kind"] == "paragraphs"
                                      else [cell for row in block["rows"] for cell in row]))

    def test_two_tables_keep_exact_body_interleaving_whitespace_and_distinct_widths(self):
        layout = self.two_table_layout()
        write_english_ocr_docx(str(self.source), str(self.output), [self.layout_text(layout), "Last page reference"],
                              page_layouts=[layout, None])
        document = Document(self.output)
        self.assertEqual(len(document.tables), 2)
        self.assertEqual([[[cell.text for cell in row.cells] for row in table.rows] for table in document.tables],
                         [layout["blocks"][1]["rows"], layout["blocks"][3]["rows"]])
        paragraphs = [p.text for p in document.paragraphs]
        expected_lines = [line for block in layout["blocks"] if block["kind"] == "paragraphs" for line in block["lines"]]
        self.assertEqual(paragraphs[:len(expected_lines)], expected_lines)
        self.assertEqual(paragraphs[-1], "Last page reference")
        blocks = [(el.tag.rsplit("}", 1)[-1], " ".join(el.xpath(".//w:t/text()"))) for el in document.element.body]
        self.assertEqual([tag for tag, text in blocks if text], ["p", "tbl", "p", "tbl", "p", "p"])
        self.assertAlmostEqual(document.tables[0].columns[0].width.pt / document.tables[0].columns[1].width.pt, 3, places=2)
        self.assertAlmostEqual(document.tables[1].columns[1].width.pt / document.tables[1].columns[0].width.pt, 2.5, places=2)
        self.assertEqual([(s.page_width.pt, s.page_height.pt) for s in document.sections], [(612, 792), (792, 612)])
        self.assertEqual(len(document.inline_shapes), 0)

    def test_adjacent_tables_have_an_empty_word_paragraph_separator(self):
        layout = self.two_table_layout()
        layout["blocks"][2]["lines"] = []
        write_english_ocr_docx(str(self.source), str(self.output), [self.layout_text(layout), "Last page"],
                              page_layouts=[layout, None])
        children = list(Document(self.output).element.body)
        indices = [index for index, node in enumerate(children) if node.tag.endswith("}tbl")]
        self.assertEqual(len(indices), 2)
        self.assertEqual(indices[1] - indices[0], 2)
        self.assertTrue(children[indices[0] + 1].tag.endswith("}p"))
        self.assertEqual(children[indices[0] + 1].xpath(".//w:t/text()"), [])

    def test_invalid_later_block_is_rejected_before_saving_partial_content(self):
        original = self.two_table_layout()
        mutations = []
        for widths in ([0.5], [0, 0.5, 0.5], [float("nan"), 0.5, 0.5], [0.3, 0.3, 0.3]):
            layout = copy.deepcopy(original)
            layout["blocks"][3]["columnWidths"] = widths
            mutations.append(layout)
        for bad_block in ({"kind": "image"}, {"kind": "table", "rows": [], "columnWidths": [0.5, 0.5]},
                          {"kind": "paragraphs", "lines": "Missing paragraph list"}):
            layout = copy.deepcopy(original)
            layout["blocks"][3] = bad_block
            mutations.append(layout)
        mutations.append({"blocks": []})
        for layout in mutations:
            with self.subTest(layout=layout), self.assertRaises(ValueError):
                write_english_ocr_docx(str(self.source), str(self.output), [self.layout_text(original), "Last page"],
                                      page_layouts=[layout, None])
            self.assertFalse(self.output.exists())

    def test_layout_cannot_omit_duplicate_or_change_a_recognized_amount(self):
        for change in ("remove", "duplicate", "amount"):
            original = self.two_table_layout()
            layout = copy.deepcopy(original)
            if change == "remove":
                layout["blocks"].pop(2)
            elif change == "duplicate":
                layout["blocks"].append(copy.deepcopy(layout["blocks"][0]))
            else:
                layout["blocks"][3]["rows"][1][2] = "120.26"
            with self.subTest(change=change), self.assertRaisesRegex(ValueError, "preserve every recognized word"):
                write_english_ocr_docx(str(self.source), str(self.output), [self.layout_text(original), "Last page"],
                                      page_layouts=[layout, None])
            self.assertFalse(self.output.exists())

    def test_txt_and_tsv_share_one_process_and_stale_second_tsv_is_not_reused(self):
        calls = []
        def run(args, **_kwargs):
            Path(args[2] + ".txt").write_text(self.raw, encoding="utf-8")
            if not calls:
                Path(args[2] + ".tsv").write_text("first page geometry", encoding="utf-8")
            calls.append(args)
            return SimpleNamespace(returncode=0, stderr="")
        layouts = ["stale caller data"]
        with patch("ocr_english_docx.shutil.which", return_value="tesseract"), \
             patch("ocr_english_docx.subprocess.run", side_effect=run), \
             patch("ocr_ruled_tables.build_ruled_page_layout", return_value=self.layout) as geometry:
            transcripts = recognize_ocr_transcripts(str(self.source), language="eng", dpi=220,
                                                    page_layouts=layouts)
        self.assertEqual(transcripts, [self.raw, self.raw])
        self.assertEqual(len(calls), 2)
        self.assertTrue(all(args[-2:] == ["txt", "tsv"] for args in calls))
        geometry.assert_called_once()
        self.assertEqual(layouts, [self.layout, None])

    def test_corrupt_optional_tsv_keeps_actual_text(self):
        def run(args, **_kwargs):
            Path(args[2] + ".txt").write_text(self.raw, encoding="utf-8")
            Path(args[2] + ".tsv").write_bytes(b"\xff\xfeinvalid")
            return SimpleNamespace(returncode=0, stderr="")
        layouts = []
        with patch("ocr_english_docx.shutil.which", return_value="tesseract"), \
             patch("ocr_english_docx.subprocess.run", side_effect=run):
            transcripts = recognize_ocr_transcripts(str(self.source), language="eng", dpi=220,
                                                    page_layouts=layouts)
        self.assertEqual(transcripts, [self.raw, self.raw])
        self.assertEqual(layouts, [None, None])


if __name__ == "__main__":
    unittest.main()
