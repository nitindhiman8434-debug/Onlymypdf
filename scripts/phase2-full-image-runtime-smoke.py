"""Exercise Word reference and scan OCR with the full Linux web image tools."""
from __future__ import annotations

import argparse
import os
import subprocess
import sys
import tempfile
from pathlib import Path

import fitz
from docx import Document
from PIL import Image, ImageDraw, ImageFont


ROOT = Path(__file__).resolve().parents[1]
CONVERTER = ROOT / "scripts" / "pdf-to-docx.py"


def convert(source: Path, output: Path, *, reference: bool = False) -> None:
    args = [sys.executable, str(CONVERTER)]
    if reference:
        args.append("--reference-transcript")
    args.extend([str(source), str(output)])
    result = subprocess.run(args, capture_output=True, text=True, timeout=120, env={
        **os.environ,
        "PDF_OCR_ENABLED": "true",
        "PDF_OCR_REQUIRED": "true",
        "PDF_OCR_LANGUAGES": "eng",
        "OMP_THREAD_LIMIT": "1",
    })
    if result.returncode:
        raise AssertionError(f"Word conversion failed: {result.stderr[-1200:]}")


def docx_text(path: Path) -> str:
    document = Document(path)
    return " ".join(paragraph.text for paragraph in document.paragraphs) + " " + " ".join(
        paragraph.text for table in document.tables for row in table.rows
        for cell in row.cells for paragraph in cell.paragraphs
    )


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--skip-ocr", action="store_true")
    options = parser.parse_args()
    with tempfile.TemporaryDirectory(prefix="phase2-image-smoke-") as temp:
        folder = Path(temp)
        source = folder / "text.pdf"
        pdf = fitz.open()
        page = pdf.new_page()
        for index in range(14):
            page.insert_text((60, 75 + index * 30), f"ONLYMYPDF EDITABLE REFERENCE LINE {index:02d} VERIFIED DOCUMENT TEXT", fontsize=13)
        pdf.save(source)
        pdf.close()
        reference = folder / "reference.docx"
        convert(source, reference, reference=True)
        assert "ONLYMYPDF EDITABLE REFERENCE LINE 13" in docx_text(reference)
        print(f"Reference transcript: {reference.stat().st_size} bytes, editable text verified")

        if options.skip_ocr:
            return
        languages = subprocess.check_output(["tesseract", "--list-langs"], text=True, stderr=subprocess.STDOUT)
        assert "eng" in languages.split() and "hin" in languages.split(), languages
        image = Image.new("RGB", (1600, 900), "white")
        draw = ImageDraw.Draw(image)
        font = ImageFont.truetype("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf", 54)
        for index in range(5):
            draw.text((90, 95 + index * 135), "ONLYMYPDF OCR TEST 2026", fill="black", font=font)
        image_path = folder / "scan.png"
        image.save(image_path)
        scan = folder / "scan.pdf"
        pdf = fitz.open()
        page = pdf.new_page(width=800, height=450)
        page.insert_image(page.rect, filename=str(image_path))
        pdf.save(scan)
        pdf.close()
        ocr_output = folder / "ocr.docx"
        convert(scan, ocr_output)
        assert "ONLYMYPDF" in docx_text(ocr_output).upper()
        print(f"Scan OCR: {ocr_output.stat().st_size} bytes, editable text verified")


if __name__ == "__main__":
    main()
