#!/usr/bin/env python3
"""English-only local PDF-to-Word content, readability and controlled-layout gate."""
from __future__ import annotations

import argparse
import hashlib
import importlib.metadata
import importlib.util
import io
import json
import os
import platform
import re
import shutil
import subprocess
import time
import unicodedata
import uuid
import zipfile
from collections import Counter
from datetime import datetime, timezone
from pathlib import Path
from xml.etree import ElementTree as ET

ROOT = Path(__file__).resolve().parents[1]
ARTIFACTS = ROOT / "tmp" / "pdfs" / "phase3.2d-english-word"
REPORT = ROOT / "quality" / "phase3-english-word" / "latest-report.json"
W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
WP = "{http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing}"
DOCX_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
spec = importlib.util.spec_from_file_location("english_word_local_transport", Path(__file__).with_name("phase3-document-corpus.py"))
transport_module = importlib.util.module_from_spec(spec)
spec.loader.exec_module(transport_module)
LIMITATIONS = [
    "Four controlled synthetic English documents; not a universal PDF-to-Word or OCR accuracy score.",
    "Only the shared loopback HTTP transport is reused. No previous language corpus is generated or run.",
    "Exact NFC comparisons collapse existing whitespace to one space; they never delete word spaces or forgive character/amount differences.",
    "The controlled originals fit comfortably on their declared pages. Equal rendered page counts are required for these cases, not asserted for arbitrary documents.",
    "Invoice row association is required. Native editable table cells are measured separately, not implied by a correct plain-text row.",
    "The sources contain no intended page pictures or multiple text columns. Unexpected page images and image-plus-text side-by-side wrappers fail this scoped layout contract.",
    "Font sizes, page dimensions and image geometry are reported; exact fonts, pixel-identical positioning, arbitrary layout fidelity and accessibility are not guaranteed.",
    "Independent LibreOffice rendering checks visible PDF text and pagination; screenshots still require visual inspection.",
    "No paid services, downloads, public deployment or non-loopback requests are performed by this runner.",
]


def normalize(value: str) -> str:
    return re.sub(r"\s+", " ", unicodedata.normalize("NFC", value)).strip()


def relative(path: Path) -> str:
    return path.resolve().relative_to(ROOT.resolve()).as_posix()


def digest(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def p(text: str, heading: bool = False) -> dict:
    return {"kind": "heading" if heading else "paragraph", "text": text}


def page_text(page: list[dict]) -> list[str]:
    result = []
    for block in page:
        result.extend(" ".join(row) for row in block["rows"]) if block["kind"] == "table" else result.append(block["text"])
    return result


def definitions() -> list[dict]:
    memo = [
        p("English document quality check", True),
        p("Reference 006731"),
        p("An editable document must keep spaces between words. The amber folder belongs to the design team, while the blue folder belongs to the review team."),
        p("Each paragraph has a different purpose. Reviewers must read the project notes before approving the final schedule."),
        p("The approved amount is 1245.60 USD. The final reference is 71936."),
    ]
    rows = [
        ["Description", "Quantity", "Unit price", "Line total"],
        ["Analysis package", "2", "120.25", "240.50"],
        ["Review service", "3", "45.10", "135.30"],
        ["Archive support", "1", "19.95", "19.95"],
    ]
    invoice = [
        p("English service invoice", True), p("Invoice number 004218"),
        p("Customer: Example Research Company"),
        p("This invoice covers the completed analysis and review services. All line items and values are synthetic test information."),
        {"kind": "table", "rows": rows}, p("Amount due: 395.75 USD"),
        p("Please keep each description with its quantity and prices. The final invoice reference is 86274."),
    ]
    first = [
        p("Project review: first page", True), p("First page reference 58324"),
        p("The research team completed the initial review on Monday. The complete first paragraph must remain readable and editable in the exported document."),
        p("The approved first page amount is 240.00 USD. Reviewers should keep the source notes with the supporting schedule."),
        p("This is the last paragraph of the first page. Its closing reference is 63128."),
    ]
    second = [
        p("Project review: second page", True), p("Second page reference 94716"),
        p("The operations team completed the final review on Tuesday. This paragraph belongs after every paragraph from the first page."),
        p("The approved second page amount is 980.75 USD. No line from this page may be silently dropped or moved to the first page."),
        p("This is the final paragraph of the whole document. Its closing reference is 25869."),
    ]
    return [
        {"caseId": "english-memo-selectable", "inputKind": "selectable", "pages": [memo], "numericValues": ["006731", "1245.60", "71936"]},
        {"caseId": "english-memo-scan", "inputKind": "scan", "pages": [memo], "numericValues": ["006731", "1245.60", "71936"]},
        {"caseId": "english-invoice-scan", "inputKind": "scan", "pages": [invoice],
         "numericValues": ["004218", "120.25", "240.50", "45.10", "135.30", "19.95", "19.95", "395.75", "86274"], "invoiceRows": rows[1:]},
        {"caseId": "english-two-page-scan", "inputKind": "scan", "pages": [first, second],
         "pageSizesPoints": [[612, 792], [792, 612]],
         "numericValues": ["58324", "240.00", "63128", "94716", "980.75", "25869"]},
    ]


def phrase_positions(text: str, expected: str) -> list[int]:
    haystack, needle = normalize(text), normalize(expected)
    if not needle:
        raise ValueError("An empty phrase cannot be scored")
    # English word boundaries prevent 'paid' from passing inside 'unpaid'.
    expression = r"(?<![A-Za-z0-9_])" + re.escape(needle) + r"(?![A-Za-z0-9_])"
    return [match.start() for match in re.finditer(expression, haystack)]


def content_checks(text: str, case: dict) -> dict:
    blocks = [text for page in case["pages"] for text in page_text(page)]
    counts = []
    for block, count in Counter(blocks).items():
        positions = phrase_positions(text, block)
        counts.append({"text": block, "expected": count, "actual": len(positions), "positions": positions})
    normalized = normalize(text)
    cursor, ordered = 0, []
    for block in blocks:
        candidates = [position for position in phrase_positions(text, block) if position >= cursor]
        position = candidates[0] if candidates else None
        ordered.append({"text": block, "position": position, "passed": position is not None})
        if position is not None:
            cursor = position + len(normalize(block))
    numbers = []
    for value, expected in Counter(case.get("numericValues", [])).items():
        # A sentence's final period is allowed; a decimal/grouping continuation is not.
        pattern = r"(?<![\d.,+\-\u2212])" + re.escape(value) + r"(?!\d|[.,]\d)"
        actual = len(re.findall(pattern, unicodedata.normalize("NFC", text)))
        numbers.append({"value": value, "expected": expected, "actual": actual, "passed": actual == expected})
    invoice = []
    for row in case.get("invoiceRows", []):
        expected = " ".join(row)
        positions = phrase_positions(text, expected)
        invoice.append({"expectedCells": row, "occurrences": len(positions), "passed": len(positions) == 1})
    exact = normalized == normalize("\n".join(blocks))
    passed = exact and all(row["actual"] == row["expected"] for row in counts) and all(row["passed"] for row in ordered + numbers + invoice)
    return {"passed": passed, "fullNormalizedTextMatches": exact, "blockOccurrences": counts,
            "orderedBlocks": ordered, "numericValues": numbers, "invoiceRowAssociation": invoice,
            "visibleText": text, "normalizedText": normalized}


def truthy(element: ET.Element | None) -> bool:
    return element is not None and element.get(W + "val", "1").lower() not in ("0", "false", "off")


def visible_paragraph(paragraph: ET.Element) -> str:
    def visit(node: ET.Element) -> str:
        if node.tag in (W + "del", W + "drawing", W + "pict", W + "instrText"):
            return ""
        if node.tag == W + "r":
            properties = node.find(W + "rPr")
            if properties is not None and any(truthy(properties.find(W + name)) for name in ("vanish", "webHidden")):
                return ""
        if node.tag == W + "t":
            return node.text or ""
        if node.tag in (W + "tab", W + "br", W + "cr"):
            return " "
        return "".join(visit(child) for child in node)
    return visit(paragraph)


def inspect_docx(body: bytes, case: dict) -> dict:
    from PIL import Image
    if not body.startswith(b"PK"):
        raise ValueError("DOCX response is not ZIP data")
    with zipfile.ZipFile(io.BytesIO(body)) as archive:
        if archive.testzip() is not None:
            raise ValueError("DOCX ZIP CRC failed")
        for name in ("[Content_Types].xml", "word/document.xml"):
            if name not in archive.namelist():
                raise ValueError("Missing required DOCX XML: " + name)
        document = ET.fromstring(archive.read("word/document.xml"))
        root = document.find(W + "body")
        if root is None:
            raise ValueError("Missing DOCX document body")
        paragraphs, native_rows = [], []

        def walk(node: ET.Element) -> None:
            if node.tag in (W + "del", W + "drawing", W + "pict"):
                return
            if node.tag == W + "p":
                value = visible_paragraph(node)
                if value.strip():
                    paragraphs.append(value)
                return
            if node.tag == W + "tr":
                native_rows.append([normalize(" ".join(visible_paragraph(p) for p in cell.iter(W + "p"))) for cell in node.findall(W + "tc")])
            for child in node:
                walk(child)

        walk(root)
        text = "\n".join(paragraphs)
        media = []
        for name in archive.namelist():
            if not name.startswith("word/media/"):
                continue
            raw = archive.read(name)
            entry = {"name": name, "bytes": len(raw), "sha256": hashlib.sha256(raw).hexdigest()}
            try:
                with Image.open(io.BytesIO(raw)) as image:
                    width, height = image.size
                    entry.update({"widthPixels": width, "heightPixels": height,
                                  "pageAspectRatio": min(abs(width / height - 612 / 792), abs(width / height - 595 / 842)) < 0.04,
                                  "largeRaster": max(width, height) >= 1000})
            except Exception:
                entry["imageDimensionsAvailable"] = False
            media.append(entry)
        wrappers = 0
        for row in root.iter(W + "tr"):
            cells = row.findall(W + "tc")
            image_cells = [bool(list(cell.iter(W + "drawing")) or list(cell.iter(W + "pict"))) for cell in cells]
            text_cells = [bool(normalize(" ".join(visible_paragraph(p) for p in cell.iter(W + "p")))) for cell in cells]
            if any(image_cells) and any(has_text and not has_image for has_text, has_image in zip(text_cells, image_cells)):
                wrappers += 1
        columns = [int(node.get(W + "num", "1")) for node in root.iter(W + "cols")]
        font_sizes = sorted({int(node.get(W + "val")) / 2 for node in root.iter(W + "sz") if (node.get(W + "val") or "").isdigit()})
        geometry = [{"widthEmu": int(node.get("cx", "0")), "heightEmu": int(node.get("cy", "0"))} for node in root.iter(WP + "extent")]
        native_table_checks = [{"expectedCells": expected,
                                "matchedRows": sum([normalize(cell) for cell in row] == expected for row in native_rows)}
                               for expected in case.get("invoiceRows", [])]
        # These fixtures contain only words and table rules, never source photos.
        layout = {"passed": not media and wrappers == 0 and max(columns, default=1) == 1,
                  "unexpectedEmbeddedImages": media, "imageTextSideBySideWrappers": wrappers,
                  "sectionColumns": columns or [1], "explicitFontSizesPt": font_sizes,
                  "imageGeometry": geometry, "nativeTableRows": native_rows,
                  "nativeInvoiceTableChecks": native_table_checks,
                  "nativeInvoiceTableReconstructed": bool(native_table_checks) and all(row["matchedRows"] == 1 for row in native_table_checks),
                  "nativeTableReconstructionRequired": False}
        content = content_checks(text, case)
        return {"passed": content["passed"] and layout["passed"], "validZip": True,
                "visibleEditableCharacters": len(text.strip()), "visibleParagraphCount": len(paragraphs),
                "content": content, "layout": layout}


def rendered_checks(page_texts: list[str], case: dict, dimensions: list[dict] | None = None) -> dict:
    expected_pages = len(case["pages"])
    per_page = []
    for index, expected_page in enumerate(case["pages"]):
        text = page_texts[index] if index < len(page_texts) else ""
        expected = normalize("\n".join(page_text(expected_page)))
        actual = normalize(text)
        per_page.append({"page": index + 1, "fullTextMatches": actual == expected,
                         "expectedText": expected, "actualText": actual})
    geometry = []
    for index, expected in enumerate(case.get("pageSizesPoints", [[612, 792]] * expected_pages)):
        actual = dimensions[index] if dimensions is not None and index < len(dimensions) else None
        geometry.append({"page": index + 1, "expectedWidthPoints": expected[0], "expectedHeightPoints": expected[1],
                         "actual": actual, "tolerancePoints": 1,
                         "passed": actual is not None and abs(actual["widthPoints"] - expected[0]) <= 1 and abs(actual["heightPoints"] - expected[1]) <= 1})
    return {"passed": len(page_texts) == expected_pages and all(item["fullTextMatches"] for item in per_page) and all(item["passed"] for item in geometry),
            "expectedPageCount": expected_pages, "actualPageCount": len(page_texts),
            "pageCountPassed": len(page_texts) == expected_pages, "perPage": per_page, "pageGeometry": geometry}


def find_office(explicit: str | None) -> str:
    for value in (explicit, os.environ.get("LIBREOFFICE_PATH"), shutil.which("soffice"), shutil.which("libreoffice"),
                  "C:/Program Files/LibreOffice/program/soffice.exe"):
        if value and Path(value).is_file():
            return str(Path(value).resolve())
    raise FileNotFoundError("An existing LibreOffice installation is required; this runner installs nothing")


def latin_font() -> str:
    if Path("C:/Windows/Fonts/arial.ttf").is_file():
        return "Arial"
    if Path("/usr/share/fonts/truetype/noto/NotoSans-Regular.ttf").is_file():
        return "Noto Sans"
    if Path("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf").is_file():
        return "DejaVu Sans"
    raise FileNotFoundError("A local Arial, Noto Sans or DejaVu Sans font is required")


def render_docx(path: Path, directory: Path, office: str, profile: str) -> Path:
    directory.mkdir(parents=True, exist_ok=True)
    destination = directory / (path.stem + ".pdf")
    # Delete only this known generated file so a stale PDF cannot satisfy a failed render.
    destination.resolve().relative_to(ARTIFACTS.resolve())
    if destination.is_file():
        destination.unlink()
    command = [office, "-env:UserInstallation=" + (ARTIFACTS / profile).resolve().as_uri(), "--headless",
               "--convert-to", "pdf", "--outdir", str(directory), str(path)]
    result = subprocess.run(command, capture_output=True, text=True, encoding="utf-8", errors="replace", timeout=90, check=False)
    if result.returncode != 0 or not destination.is_file():
        raise ValueError("LibreOffice failed to render the DOCX")
    return destination


def inspect_render(path: Path, case: dict, prefix: str) -> dict:
    import fitz
    with fitz.open(path) as pdf:
        text = [page.get_text(sort=False) for page in pdf]
        dimensions = [{"widthPoints": page.rect.width, "heightPoints": page.rect.height} for page in pdf]
        result = rendered_checks(text, case, dimensions)
        previews = []
        for index, page in enumerate(pdf, 1):
            preview = ARTIFACTS / f"{prefix}-page-{index}.png"
            page.get_pixmap(dpi=120, colorspace=fitz.csRGB, alpha=False).save(preview)
            previews.append(relative(preview))
        return {**result, "pdfPath": relative(path), "pdfSha256": digest(path), "previewPaths": previews,
                "pageDimensions": dimensions}


def generate_source(case: dict, office: str) -> dict:
    import fitz
    from docx import Document
    from docx.enum.section import WD_SECTION_START
    from docx.shared import Inches, Pt
    ARTIFACTS.mkdir(parents=True, exist_ok=True)
    source = ARTIFACTS / (case["caseId"] + "-original.docx")
    document = Document()
    font = latin_font()
    normal = document.styles["Normal"]
    normal.font.name, normal.font.size = font, Pt(12)
    normal.paragraph_format.space_after, normal.paragraph_format.line_spacing = Pt(10), 1.1
    for section in document.sections:
        section.page_width, section.page_height = Inches(8.5), Inches(11)
        section.left_margin = section.right_margin = Inches(0.7)
        section.top_margin = section.bottom_margin = Inches(0.7)
    for page_index, page in enumerate(case["pages"]):
        if page_index:
            document.add_section(WD_SECTION_START.NEW_PAGE)
        width, height = case.get("pageSizesPoints", [[612, 792]] * len(case["pages"]))[page_index]
        document.sections[-1].page_width, document.sections[-1].page_height = Pt(width), Pt(height)
        for block in page:
            if block["kind"] == "table":
                table = document.add_table(rows=0, cols=len(block["rows"][0]))
                table.style = "Table Grid"
                for values in block["rows"]:
                    cells = table.add_row().cells
                    for cell, value in zip(cells, values):
                        cell.text = value
                document.add_paragraph()
            else:
                run = document.add_paragraph().add_run(block["text"])
                if block["kind"] == "heading":
                    run.bold, run.font.size = True, Pt(17)
    document.save(source)
    original_check = inspect_docx(source.read_bytes(), case)
    if not original_check["content"]["passed"]:
        raise ValueError("Literal source DOCX content validation failed")
    original_pdf = render_docx(source, ARTIFACTS / "original-pdfs", office, "source-libreoffice-profile")
    source_render = inspect_render(original_pdf, case, case["caseId"] + "-source-preview")
    if not source_render["passed"]:
        raise ValueError("Original PDF source text or pagination failed; conversion cannot be scored")
    upload = ARTIFACTS / (case["caseId"] + ".pdf")
    if case["inputKind"] == "selectable":
        shutil.copyfile(original_pdf, upload)
    else:
        with fitz.open(original_pdf) as original, fitz.open() as scan:
            for index, page in enumerate(original, 1):
                image = page.get_pixmap(dpi=300, colorspace=fitz.csRGB, alpha=False)
                png = ARTIFACTS / f"{case['caseId']}-upload-image-{index}.png"
                png.write_bytes(image.tobytes("png"))
                scan.new_page(width=page.rect.width, height=page.rect.height).insert_image(page.rect, filename=str(png))
            scan.save(upload, garbage=4, deflate=True)
    with fitz.open(upload) as pdf:
        counts = [len(page.get_text().strip()) for page in pdf]
        valid = len(pdf) == len(case["pages"]) and (all(count == 0 for count in counts) if case["inputKind"] == "scan" else all(count > 0 for count in counts))
    return {"passed": valid, "inputPath": relative(upload), "inputSha256": digest(upload), "inputBytes": upload.stat().st_size,
            "inputKind": case["inputKind"], "selectableCharactersByPage": counts, "fontFamily": font,
            "literalDocxPath": relative(source), "literalDocxSha256": digest(source),
            "literalDocxTextPassed": True, "trustedOriginal": source_render}


def multipart(path: Path) -> tuple[bytes, str]:
    boundary = "onlymypdf-english-word-" + uuid.uuid4().hex
    data = (f'--{boundary}\r\nContent-Disposition: form-data; name="file"; filename="{path.name}"\r\n'
            'Content-Type: application/pdf\r\n\r\n').encode() + path.read_bytes() + b"\r\n"
    data += f'--{boundary}\r\nContent-Disposition: form-data; name="options"\r\n\r\n{{}}\r\n--{boundary}--\r\n'.encode()
    return data, f"multipart/form-data; boundary={boundary}"


def convert(case: dict, base: str, timeout: float, office: str) -> dict:
    start = time.monotonic()
    result = {"caseId": case["caseId"], "tool": "pdf-to-word"}
    try:
        transport = transport_module.LocalTransport(base, timeout)
        payload, content_type = multipart(ROOT / case["source"]["inputPath"])
        status, body, headers = transport.request("/api/tools/pdf-to-word", payload, content_type)
        mime = headers.get("content-type", "").split(";", 1)[0].strip().lower()
        result.update({"httpStatus": status, "mime": mime, "outputBytes": len(body),
                       "responseSha256": hashlib.sha256(body).hexdigest(), "responseEngine": headers.get("x-pdf-engine"),
                       "conversionDurationMs": round((time.monotonic() - start) * 1000)})
        if status != 200 or mime != DOCX_MIME:
            result.update({"status": "failed", "reason": "Expected HTTP 200 with DOCX MIME"})
            return result
        output = ARTIFACTS / (case["caseId"] + "-result.docx")
        output.write_bytes(body)
        result["outputPath"] = relative(output)
        result["document"] = inspect_docx(body, case)
        # Render every successful artifact once, even if XML checks already failed.
        rendered = render_docx(output, ARTIFACTS / "rendered-outputs", office, "output-libreoffice-profile")
        result["rendered"] = inspect_render(rendered, case, case["caseId"] + "-output-preview")
        result["status"] = "passed" if result["document"]["passed"] and result["rendered"]["passed"] else "failed"
    except Exception as exc:
        result.update({"status": "error", "error": str(exc)[:500]})
    finally:
        result["totalDurationMs"] = round((time.monotonic() - start) * 1000)
    return result


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--base-url", default="http://127.0.0.1:3001")
    parser.add_argument("--generate-only", action="store_true")
    parser.add_argument("--case", action="append", choices=[case["caseId"] for case in definitions()])
    parser.add_argument("--report", type=Path, default=REPORT)
    parser.add_argument("--timeout", type=float, default=180)
    parser.add_argument("--libreoffice")
    args = parser.parse_args()
    if not 1 <= args.timeout <= 600:
        parser.error("timeout must be between 1 and 600 seconds")
    base = transport_module.validate_base_url(args.base_url)
    selected = [case for case in definitions() if not args.case or case["caseId"] in args.case]
    report = {"schemaVersion": 1, "phase": "3.2D", "generatedAt": datetime.now(timezone.utc).isoformat(),
              "baseUrl": base, "generateOnly": args.generate_only, "limitations": LIMITATIONS,
              "environment": {"python": platform.python_version(), "platform": platform.system(),
                              "packages": {name: importlib.metadata.version(name) for name in ("PyMuPDF", "python-docx", "Pillow")}},
              "cases": [], "results": []}
    try:
        office = find_office(args.libreoffice)
        for case in selected:
            item = {**case, "source": generate_source(case, office)}
            report["cases"].append(item)
            if item["source"]["passed"] and not args.generate_only:
                result = convert(item, base, args.timeout, office)
                report["results"].append(result)
                print(json.dumps({"caseId": case["caseId"], "status": result["status"], "httpStatus": result.get("httpStatus")}), flush=True)
    except Exception as exc:
        report["setupError"] = str(exc)[:500]
    fixture_failures = sum(not case["source"]["passed"] for case in report["cases"])
    statuses = Counter(item["status"] for item in report["results"])
    complete = len(report["cases"]) == len(selected) and not fixture_failures and not report.get("setupError")
    all_passed = complete and len(report["results"]) == len(selected) and statuses["passed"] == len(selected)
    report["summary"] = {"casesSelected": len(selected), "fixturesGenerated": len(report["cases"]),
                         "fixturesFailed": fixture_failures, "routesTested": len(report["results"]),
                         "passed": statuses["passed"], "failed": statuses["failed"], "errors": statuses["error"],
                         "allPassed": None if args.generate_only else bool(all_passed)}
    args.report.parent.mkdir(parents=True, exist_ok=True)
    args.report.write_text(json.dumps(report, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps({"report": str(args.report), "summary": report["summary"]}), flush=True)
    return 0 if (complete if args.generate_only else all_passed) else 1


if __name__ == "__main__":
    raise SystemExit(main())
