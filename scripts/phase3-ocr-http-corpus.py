#!/usr/bin/env python3
"""Strict synthetic OCR checks against local HTTP routes; no remote services."""
from __future__ import annotations

import argparse
import hashlib
import importlib.metadata
import importlib.util
import io
import json
import os
import platform
import re
import shutil
import subprocess
import sys
import time
import unicodedata
import uuid
import zipfile
from datetime import datetime, timezone
from pathlib import Path
from xml.etree import ElementTree as ET

ROOT = Path(__file__).resolve().parents[1]
ARTIFACTS = ROOT / "tmp" / "pdfs" / "phase3.2c-ocr"
REPORT = ROOT / "quality" / "phase3-ocr" / "latest-report.json"
spec = importlib.util.spec_from_file_location("document_corpus_transport", Path(__file__).with_name("phase3-document-corpus.py"))
shared = importlib.util.module_from_spec(spec)
spec.loader.exec_module(shared)
W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
DOCX_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
MIMES = {"ocr-pdf": "application/pdf", "pdf-to-word": DOCX_MIME}
LIMITATIONS = [
    "Synthetic English and Hindi scans on one local machine, not arbitrary scripts, customer evidence or universal OCR accuracy.",
    "Trusted originals are DOCX generated from literal Unicode and rendered by LibreOffice; PyMuPDF validates source ActualText before rasterization. LibreOffice Hindi glyph mappings are not portable across all PDF text extractors.",
    "Image-only inputs must contain zero selectable text on every page. No source text or expected answers are sent to the conversion routes.",
    "Selected text markers require exact Unicode NFC text after whitespace removal; numeric markers require complete contiguous values with numeric boundaries. There is no fuzzy matching, transliteration, dropped accents or numerical tolerance.",
    "Selected text and amount checks do not establish complete reading order, document layout, OCR confidence calibration or accessibility semantics.",
    "Source and output previews require human visual review; the runner does not infer visual correctness from text extraction.",
    "Runner process OCR environment and models are recorded separately from response engine headers; they do not prove the server used the same environment.",
    "No paid APIs, model downloads, cloud deployment, external URLs or environment HTTP proxies are used by this runner.",
]


def stamp() -> str:
    return datetime.now(timezone.utc).isoformat()


def relative(path: Path) -> str:
    return path.resolve().relative_to(ROOT).as_posix()


def sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def marker_check(text: str, expected: list[dict]) -> dict:
    """Count exact selected text; a numeric marker cannot be part of another value."""
    normalized = shared.normalize_text(text)
    numeric_text = unicodedata.normalize("NFC", text)
    counts = []
    for marker in expected:
        needle = shared.normalize_text(marker["text"])
        if not needle:
            raise ValueError("Empty normalized marker is invalid")
        if re.fullmatch(r"[+-]?\d+(?:[.,]\d+)*", needle):
            # Preserve numeric meaning: 12345.67 cannot match 912345.67,
            # 12345.670, -12345.67 or a component of a grouped decimal.
            pattern = r"(?<![\d.,+\-\u2212])" + re.escape(needle) + r"(?![\d.,])"
            # Keep whitespace as a boundary so two separate identical amounts
            # count twice rather than becoming one larger concatenated number.
            actual = len(re.findall(pattern, numeric_text))
        else:
            actual = normalized.count(needle)
        counts.append({"text": marker["text"], "expected": marker["expected"], "actual": actual})
    missing = [row for row in counts if row["actual"] < row["expected"]]
    duplicated = [row for row in counts if row["actual"] > row["expected"]]
    return {"passed": not missing and not duplicated, "missing": missing,
            "duplicated": duplicated, "counts": counts}


def definitions() -> list[dict]:
    english = [
        "OnlyMyPDF scan quality check",
        "Invoice 58324",
        "This scanned page must become searchable and editable.",
        "Customer: Example Research Company",
        "Total amount 1475.25",
        "Final reference 71936",
    ]
    hindi = [
        "देवनागरी संयुक्त अक्षर परीक्षण",
        "क्षेत्र त्रिकोण ज्ञान श्रद्धा",
        "प्रक्रिया स्वतंत्र राष्ट्रीय दृष्टि",
        "उत्कृष्ट विश्वविद्यालय",
        "यह स्कैन संपादन योग्य पाठ में बदलना चाहिए।",
        "Invoice 48219",
        "कुल राशि 12345.67 रुपये।",
        "Final reference 86427",
    ]
    return [
        {"caseId": "clean-english", "language": "eng", "kind": "scan", "lines": english,
         "markers": shared.markers("58324", "searchable and editable", "Example Research Company", "1475.25", "71936"),
         "tools": ["ocr-pdf", "pdf-to-word"], "expectedStatus": 200},
        {"caseId": "hindi-conjunct-bilingual", "language": "eng+hin", "kind": "scan", "lines": hindi,
         "markers": shared.markers("क्षेत्र", "त्रिकोण", "ज्ञान", "श्रद्धा", "प्रक्रिया", "स्वतंत्र", "राष्ट्रीय", "दृष्टि", "उत्कृष्ट", "विश्वविद्यालय", "48219", "कुल राशि", "12345.67", "86427"),
         "tools": ["ocr-pdf", "pdf-to-word"], "expectedStatus": 200},
        {"caseId": "existing-searchable", "language": "eng", "kind": "searchable", "fromCase": "clean-english",
         "markers": shared.markers("58324", "searchable and editable", "1475.25", "71936"),
         "tools": ["ocr-pdf"], "expectedStatus": 200, "expectedText": "\n".join(english)},
        {"caseId": "unreadable-raster", "language": "eng+hin", "kind": "unreadable", "markers": [],
         "tools": ["ocr-pdf", "pdf-to-word"], "expectedStatus": 422},
        {"caseId": "searchable-plus-unreadable", "language": "eng+hin", "kind": "mixed-unreadable", "markers": [],
         "tools": ["ocr-pdf"], "expectedStatus": 422},
    ]


def find_office(explicit: str | None) -> str:
    candidates = [explicit, os.environ.get("LIBREOFFICE_PATH"), shutil.which("soffice"), shutil.which("libreoffice"),
                  "C:/Program Files/LibreOffice/program/soffice.exe"]
    for value in candidates:
        if value and Path(value).is_file():
            return str(Path(value).resolve())
    raise FileNotFoundError("LibreOffice is required to generate independently rendered OCR sources; no installer is run.")


def font_names() -> tuple[str, str]:
    if Path("C:/Windows/Fonts/Nirmala.ttc").is_file():
        return "Arial", "Nirmala UI"
    if any(Path(path).is_file() for path in (
            "/usr/share/fonts/truetype/noto/NotoSansDevanagari-Regular.ttf",
            "/usr/share/fonts/opentype/noto/NotoSansDevanagari-Regular.ttf")):
        return "Noto Sans", "Noto Sans Devanagari"
    raise FileNotFoundError("A local Nirmala UI or Noto Sans Devanagari font is required; no font is downloaded.")


def docx_text(data: bytes) -> tuple[str, int]:
    if not data.startswith(b"PK"):
        raise ValueError("DOCX response is not ZIP data")
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        if archive.testzip() is not None:
            raise ValueError("DOCX ZIP CRC failed")
        for required in ("[Content_Types].xml", "word/document.xml"):
            if required not in archive.namelist():
                raise ValueError("Missing required DOCX XML: " + required)
        root = ET.fromstring(archive.read("word/document.xml"))
        paragraphs = ["".join(node.text or "" for node in p.iter(W + "t")) for p in root.iter(W + "p")]
        # Only editable document text counts, never image alt text, comments or metadata.
        return "\n".join(paragraphs), sum(1 for name in archive.namelist() if name.startswith("word/media/"))


def generate_trusted(case: dict, directory: Path, office: str, fonts: tuple[str, str]) -> dict:
    import fitz
    from docx import Document
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Inches, Pt

    original = directory / (case["caseId"] + "-original.docx")
    document = Document()
    section = document.sections[0]
    section.page_width, section.page_height = Inches(8.5), Inches(11)
    section.left_margin = section.right_margin = Inches(0.7)
    section.top_margin = section.bottom_margin = Inches(0.65)
    for line in case["lines"]:
        paragraph = document.add_paragraph()
        paragraph.paragraph_format.space_after = Pt(12)
        run = paragraph.add_run(line)
        family = fonts[1] if re.search(r"[\u0900-\u097f]", line) else fonts[0]
        run.font.name, run.font.size = family, Pt(20)
        properties = run._element.get_or_add_rPr()
        properties.rFonts.set(qn("w:cs"), fonts[1])
        properties.rFonts.set(qn("w:eastAsia"), family)
        size = OxmlElement("w:szCs")
        size.set(qn("w:val"), "40")
        properties.append(size)
        language = OxmlElement("w:lang")
        language.set(qn("w:val"), "hi-IN" if family == fonts[1] else "en-US")
        language.set(qn("w:bidi"), "hi-IN")
        properties.append(language)
    document.save(original)
    expected = "\n".join(case["lines"])
    actual_docx, _ = docx_text(original.read_bytes())
    if actual_docx != expected:
        raise ValueError("Original DOCX Unicode differs from literal fixture source")
    profile = (directory / "libreoffice-profile").resolve().as_uri()
    command = [office, "-env:UserInstallation=" + profile, "--headless", "--convert-to", "pdf",
               "--outdir", str(directory), str(original)]
    completed = subprocess.run(command, capture_output=True, text=True, encoding="utf-8", errors="replace", timeout=90, check=False)
    pdf_path = original.with_suffix(".pdf")
    if completed.returncode != 0 or not pdf_path.is_file():
        raise ValueError("LibreOffice source rendering failed")
    previews, images = [], []
    with fitz.open(pdf_path) as pdf:
        text = "\n".join(page.get_text(sort=False) for page in pdf)
        if len(pdf) != 1 or shared.normalize_text(text) != shared.normalize_text(expected):
            raise ValueError("Trusted PDF text/page gate failed; cannot score OCR against a corrupt original")
        for index, page in enumerate(pdf, 1):
            png = directory / f"{case['caseId']}-source-{index}.png"
            # The saved source image is also the precise image sent inside the PDF.
            page.get_pixmap(dpi=300, colorspace=fitz.csRGB, alpha=False).save(png)
            images.append((png, page.rect.width, page.rect.height))
            previews.append(relative(png))
        actual_text = any(b"ActualText" in (pdf.xref_stream(index) or b"") for index in range(1, pdf.xref_length()) if pdf.xref_is_stream(index))
    return {"originalDocxPath": relative(original), "originalDocxSha256": sha256(original),
            "trustedPdfPath": relative(pdf_path), "trustedPdfSha256": sha256(pdf_path),
            "sourceUnicodeExact": True, "sourceContainsActualText": actual_text,
            "sourceValidationEngine": "PyMuPDF ActualText extraction and exact original DOCX XML",
            "fontFamilies": list(fonts), "previewPaths": previews, "images": images}


def image_pdf(images: list[tuple[Path, float, float]], destination: Path) -> None:
    import fitz
    with fitz.open() as document:
        for path, width, height in images:
            page = document.new_page(width=width, height=height)
            page.insert_image(page.rect, filename=str(path))
        document.save(destination, garbage=4, deflate=True)


def source_state(path: Path, expected_pages: int, image_only: bool) -> dict:
    import fitz
    with fitz.open(path) as document:
        counts = [len(page.get_text().strip()) for page in document]
        passed = len(document) == expected_pages and (not image_only or all(count == 0 for count in counts))
        return {"path": relative(path), "sha256": sha256(path), "bytes": path.stat().st_size,
                "pages": len(document), "selectableCharactersByPage": counts,
                "requiresZeroSelectableText": image_only, "passed": passed}


def generate_sources(cases: list[dict], directory: Path, office: str) -> list[dict]:
    import fitz
    from PIL import Image, ImageDraw
    directory.mkdir(parents=True, exist_ok=True)
    generated = {}
    fonts = font_names()
    for case in cases:
        path = directory / (case["caseId"] + ".pdf")
        metadata = {}
        pages, image_only = 1, True
        if case["kind"] == "scan":
            metadata = generate_trusted(case, directory, office, fonts)
            image_pdf(metadata.pop("images"), path)
        elif case["kind"] == "searchable":
            trusted = ROOT / generated[case["fromCase"]]["source"]["trustedPdfPath"]
            shutil.copyfile(trusted, path)
            image_only = False
        elif case["kind"] == "unreadable":
            image = Image.new("RGB", (1275, 1650), "white")
            draw = ImageDraw.Draw(image)
            # Smooth geometric shapes contain no letters or digits to recognize.
            draw.ellipse((120, 190, 1110, 1180), fill="#d8dce1")
            draw.rectangle((250, 1280, 1020, 1330), fill="#b9c1ca")
            png = directory / "unreadable-raster-source.png"
            image.save(png)
            image_pdf([(png, 612, 792)], path)
            metadata["previewPaths"] = [relative(png)]
        else:
            with fitz.open() as document:
                for source_name in ("existing-searchable", "unreadable-raster"):
                    with fitz.open(ROOT / generated[source_name]["source"]["path"]) as part:
                        document.insert_pdf(part)
                document.save(path, garbage=4, deflate=True)
            pages, image_only = 2, False
        state = source_state(path, pages, image_only)
        if case["kind"] == "mixed-unreadable":
            state["passed"] = state["passed"] and state["selectableCharactersByPage"][0] > 0 and state["selectableCharactersByPage"][1] == 0
        if case["kind"] == "searchable":
            with fitz.open(path) as document:
                state["passed"] = state["passed"] and shared.normalize_text(document[0].get_text()) == shared.normalize_text(case["expectedText"])
        item = {**case, "source": {**state, **metadata}}
        generated[case["caseId"]] = item
    return list(generated.values())


def multipart(path: Path, language: str) -> tuple[bytes, str]:
    boundary = "onlymypdf-ocr-corpus-" + uuid.uuid4().hex
    payload = (f'--{boundary}\r\nContent-Disposition: form-data; name="file"; filename="{path.name}"\r\n'
               'Content-Type: application/pdf\r\n\r\n').encode() + path.read_bytes() + b"\r\n"
    for name, value in (("languages", language), ("options", "{}")):
        payload += f'--{boundary}\r\nContent-Disposition: form-data; name="{name}"\r\n\r\n{value}\r\n'.encode()
    payload += f"--{boundary}--\r\n".encode()
    return payload, f"multipart/form-data; boundary={boundary}"


def inspect_output(body: bytes, mime: str, tool: str, case: dict, output: Path) -> dict:
    expected_mime = MIMES[tool]
    if mime.split(";", 1)[0].strip().lower() != expected_mime:
        raise ValueError("Wrong response MIME type")
    if tool == "pdf-to-word":
        text, media = docx_text(body)
        result = {"validOfficeZip": True, "editableCharacters": len(text.strip()), "mediaItems": media}
    else:
        import fitz
        if not body.startswith(b"%PDF-"):
            raise ValueError("OCR response is not PDF data")
        with fitz.open(stream=body, filetype="pdf") as document:
            page_texts = [page.get_text(sort=False) for page in document]
            text = "\n".join(page_texts)
            result = {"validPdf": True, "pages": len(document), "expectedPages": case["source"]["pages"],
                      "pageCountPassed": len(document) == case["source"]["pages"],
                      "searchableCharactersByPage": [len(value.strip()) for value in page_texts],
                      "everyPageSearchable": bool(page_texts) and all(value.strip() for value in page_texts)}
            previews = []
            for index, page in enumerate(document, 1):
                preview = output.with_name(output.stem + f"-page-{index}.png")
                page.get_pixmap(dpi=120, alpha=False).save(preview)
                previews.append(relative(preview))
            result["previewPaths"] = previews
    check = marker_check(text, case["markers"])
    result["markerCheck"] = check
    result["text"] = text
    result["passed"] = check["passed"] and bool(text.strip())
    if tool == "ocr-pdf":
        result["passed"] = result["passed"] and result["pageCountPassed"] and result["everyPageSearchable"]
    if "expectedText" in case:
        result["existingTextPreservedExactly"] = shared.normalize_text(text) == shared.normalize_text(case["expectedText"])
        result["passed"] = result["passed"] and result["existingTextPreservedExactly"]
    return result


def expected_rejection(status: int, body: bytes, mime: str) -> dict:
    json_type = mime.split(";", 1)[0].strip().lower() == "application/json"
    try:
        error = json.loads(body) if json_type else None
    except (ValueError, UnicodeDecodeError):
        error = None
    message = error.get("error") if isinstance(error, dict) else None
    success_artifact = body.startswith((b"%PDF-", b"PK"))
    unreadable_reason = isinstance(message, str) and bool(re.fullmatch(
        r"PDF OCR could not (?:recognize readable text on page [1-9]\d*\. Try a clearer scan or supported language\.|"
        r"extract enough readable text\. Try a clearer scan or supported OCR language\.)", message.strip()))
    return {"expectedHttpStatus": 422, "passed": status == 422 and unreadable_reason and not success_artifact,
            "jsonError": isinstance(message, str), "errorMessage": message[:400] if isinstance(message, str) else None,
            "unreadableReasonMatched": unreadable_reason,
            "successArtifactReturned": success_artifact}


def run_case(case: dict, tool: str, base: str, timeout: float, directory: Path) -> dict:
    start = time.monotonic()
    result = {"caseId": case["caseId"], "tool": tool, "expectedHttpStatus": case["expectedStatus"]}
    try:
        transport = shared.LocalTransport(base, timeout)
        payload, content_type = multipart(ROOT / case["source"]["path"], case["language"])
        # Synchronous Word route makes real error status observable; no job error is relabeled 422.
        status, body, headers = transport.request("/api/tools/" + tool, payload, content_type)
        mime = headers.get("content-type", "")
        result.update({"httpStatus": status, "mime": mime, "outputBytes": len(body),
                       "responseSha256": hashlib.sha256(body).hexdigest(), "responseEngine": headers.get("x-pdf-engine")})
        if case["expectedStatus"] == 422:
            checks = expected_rejection(status, body, mime)
        elif status == 200:
            extension = ".pdf" if tool == "ocr-pdf" else ".docx"
            output = directory / f"{case['caseId']}-{tool}{extension}"
            output.write_bytes(body)
            result["outputPath"] = relative(output)
            checks = inspect_output(body, mime, tool, case, output)
        else:
            checks = {"passed": False, "reason": "Expected a successful conversion response"}
            rejection = expected_rejection(status, body, mime)
            if rejection["errorMessage"]:
                checks["errorMessage"] = rejection["errorMessage"]
        result["checks"] = checks
        result["status"] = "passed" if checks["passed"] else "failed"
    except Exception as exc:
        result.update({"status": "error", "error": str(exc)[:600]})
    result["durationMs"] = round((time.monotonic() - start) * 1000)
    return result


def runner_environment(office: str) -> dict:
    import fitz
    packages = {}
    for name in ("PyMuPDF", "python-docx", "Pillow"):
        try:
            packages[name] = importlib.metadata.version(name)
        except importlib.metadata.PackageNotFoundError:
            packages[name] = None
    commands = {"libreoffice": [office, "--version"]}
    tesseract = shutil.which(os.environ.get("TESSERACT_PATH", "tesseract"))
    if tesseract:
        commands["tesseract"] = [tesseract, "--version"]
        commands["availableLanguages"] = [tesseract, "--list-langs"]
    versions = {}
    for name, command in commands.items():
        try:
            run = subprocess.run(command, capture_output=True, text=True, encoding="utf-8", errors="replace", timeout=15, check=False)
            lines = (run.stdout + run.stderr).strip().splitlines()
            # --list-langs prints its absolute data path in a heading; publish only codes.
            versions[name] = ([line.strip() for line in lines if re.fullmatch(r"[a-z][a-z0-9_]{1,31}", line.strip())]
                              if name == "availableLanguages" else lines[:12])
        except (OSError, subprocess.TimeoutExpired):
            versions[name] = ["Version query unavailable"]
    model_dir = os.environ.get("TESSDATA_PREFIX")
    if not model_dir:
        try:
            model_dir = fitz.get_tessdata()
        except Exception:
            model_dir = None
    models = []
    if model_dir:
        for name in ("eng", "hin", "osd"):
            path = Path(model_dir) / (name + ".traineddata")
            if path.is_file():
                models.append({"language": name, "bytes": path.stat().st_size, "sha256": sha256(path)})
    allowlist = ("PDF_OCR_ENABLED", "PDF_OCR_REQUIRED", "PDF_OCR_LANGUAGES", "PDF_OCR_DPI", "PDF_OCR_MAX_PAGES", "OMP_THREAD_LIMIT")
    return {"python": platform.python_version(), "platform": platform.system(), "architecture": platform.machine(),
            "packages": packages, "versions": versions, "runnerOcrEnvironment": {key: os.environ.get(key) for key in allowlist},
            "runnerModels": models, "serverEnvironmentVerified": False}


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--base-url", default="http://127.0.0.1:3001")
    parser.add_argument("--generate-only", action="store_true")
    parser.add_argument("--case", action="append", choices=[case["caseId"] for case in definitions()])
    parser.add_argument("--tool", action="append", choices=list(MIMES))
    parser.add_argument("--timeout", type=float, default=180)
    parser.add_argument("--libreoffice")
    parser.add_argument("--report", type=Path, default=REPORT)
    args = parser.parse_args()
    if not 1 <= args.timeout <= 600:
        parser.error("timeout must be between 1 and 600 seconds")
    base = shared.validate_base_url(args.base_url)
    report = {"schemaVersion": 1, "phase": "3.2C", "generatedAt": stamp(), "baseUrl": base,
              "generateOnly": args.generate_only, "limitations": LIMITATIONS, "cases": [], "results": []}
    try:
        office = find_office(args.libreoffice)
        report["environment"] = runner_environment(office)
        all_cases = generate_sources(definitions(), ARTIFACTS, office)
        selected = [case for case in all_cases if not args.case or case["caseId"] in args.case]
        report["cases"] = selected
        for case in selected:
            if not case["source"]["passed"] or args.generate_only:
                continue
            for tool in case["tools"]:
                if args.tool and tool not in args.tool:
                    continue
                result = run_case(case, tool, base, args.timeout, ARTIFACTS)
                report["results"].append(result)
                print(json.dumps({"caseId": case["caseId"], "tool": tool, "status": result["status"],
                                  "durationMs": result["durationMs"], "httpStatus": result.get("httpStatus")}), flush=True)
    except Exception as exc:
        report["setupError"] = str(exc)[:600]
    failures = sum(not case["source"]["passed"] for case in report["cases"])
    passed = sum(item["status"] == "passed" for item in report["results"])
    failed = sum(item["status"] == "failed" for item in report["results"])
    errors = sum(item["status"] == "error" for item in report["results"])
    all_passed = bool(report["results"]) and not (failures or failed or errors or report.get("setupError"))
    report["summary"] = {"cases": len(report["cases"]), "fixturesPassed": len(report["cases"]) - failures,
                         "fixturesFailed": failures, "routesTested": len(report["results"]), "passed": passed,
                         "failed": failed, "errors": errors, "allPassed": None if args.generate_only else all_passed}
    args.report.parent.mkdir(parents=True, exist_ok=True)
    args.report.write_text(json.dumps(report, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps({"report": str(args.report), "summary": report["summary"]}), flush=True)
    if args.generate_only:
        return 0 if report["cases"] and not failures and not report.get("setupError") else 1
    return 0 if all_passed else 1


if __name__ == "__main__":
    raise SystemExit(main())
