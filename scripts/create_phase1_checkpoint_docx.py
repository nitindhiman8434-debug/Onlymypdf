from __future__ import annotations

import os
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_ROW_HEIGHT_RULE, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUTPUT = Path(
    os.environ.get(
        "PHASE1_DOCX_OUTPUT",
        r"C:\Users\NiTiN Dhiman\Documents\Cursor\onlymypdf-audit-2026-09-17\OnlyMyPDF_Phase_1_Dependable_Beta_Implementation_Checkpoint.docx",
    )
)

NAVY = "17365D"
BLUE = "2F75B5"
LIGHT_BLUE = "DDEBF7"
PALE_BLUE = "F2F7FC"
PALE_GRAY = "F7F7F7"
BORDER = "D9D9D9"
BLACK = RGBColor(0, 0, 0)
MUTED = RGBColor(89, 89, 89)


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=110, start=120, bottom=110, end=120) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_cell_width(cell, width_inches: float) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(int(width_inches * 1440)))
    tc_w.set(qn("w:type"), "dxa")


def set_table_borders(table, color: str = BORDER, size: str = "5") -> None:
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = borders.find(qn(f"w:{edge}"))
        if tag is None:
            tag = OxmlElement(f"w:{edge}")
            borders.append(tag)
        tag.set(qn("w:val"), "single")
        tag.set(qn("w:sz"), size)
        tag.set(qn("w:color"), color)


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    repeat = OxmlElement("w:tblHeader")
    repeat.set(qn("w:val"), "true")
    tr_pr.append(repeat)


def prevent_row_split(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def add_field(paragraph, instruction: str) -> None:
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = instruction
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, separate, text, end])


def clear_paragraph_borders(paragraph) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    borders = p_pr.find(qn("w:pBdr"))
    if borders is None:
        borders = OxmlElement("w:pBdr")
        p_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "between", "bar"):
        node = borders.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            borders.append(node)
        node.set(qn("w:val"), "nil")


def add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("Page ")
    run.font.size = Pt(9)
    run.font.color.rgb = MUTED
    add_field(paragraph, "PAGE")
    run = paragraph.add_run(" of ")
    run.font.size = Pt(9)
    run.font.color.rgb = MUTED
    add_field(paragraph, "NUMPAGES")


def add_heading(doc: Document, text: str, level: int = 1):
    paragraph = doc.add_paragraph(style=f"Heading {level}")
    paragraph.paragraph_format.keep_with_next = True
    paragraph.add_run(text)
    return paragraph


def add_body(doc: Document, text: str, bold_lead: str | None = None):
    paragraph = doc.add_paragraph(style="Body Text")
    if bold_lead and text.startswith(bold_lead):
        paragraph.add_run(bold_lead).bold = True
        paragraph.add_run(text[len(bold_lead) :])
    else:
        paragraph.add_run(text)
    return paragraph


def add_bullets(doc: Document, items: list[str]) -> None:
    for item in items:
        paragraph = doc.add_paragraph(style="List Bullet")
        paragraph.paragraph_format.space_after = Pt(4)
        paragraph.add_run(item)


def add_numbered(doc: Document, items: list[str]) -> None:
    for item in items:
        paragraph = doc.add_paragraph(style="List Number")
        paragraph.paragraph_format.space_after = Pt(5)
        paragraph.add_run(item)


def add_table(
    doc: Document,
    headers: list[str],
    rows: list[list[str]],
    widths: list[float],
    font_size: float = 9.0,
    header_fill: str = NAVY,
    center_columns: set[int] | None = None,
):
    center_columns = center_columns or set()
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_borders(table)
    table.rows[0].height_rule = WD_ROW_HEIGHT_RULE.AT_LEAST
    set_repeat_table_header(table.rows[0])
    for index, header in enumerate(headers):
        cell = table.rows[0].cells[index]
        set_cell_width(cell, widths[index])
        set_cell_shading(cell, header_fill)
        set_cell_margins(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        paragraph = cell.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.space_after = Pt(0)
        run = paragraph.add_run(header)
        run.bold = True
        run.font.size = Pt(font_size)
        run.font.color.rgb = RGBColor(255, 255, 255)

    for row_index, values in enumerate(rows):
        row = table.add_row()
        prevent_row_split(row)
        fill = "FFFFFF" if row_index % 2 == 0 else PALE_BLUE
        for column_index, value in enumerate(values):
            cell = row.cells[column_index]
            set_cell_width(cell, widths[column_index])
            set_cell_shading(cell, fill)
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.paragraph_format.line_spacing = 1.08
            if column_index in center_columns:
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = paragraph.add_run(value)
            run.font.size = Pt(font_size)
            run.font.color.rgb = BLACK

    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(1)
    return table


def add_page_break(doc: Document) -> None:
    paragraph = doc.add_paragraph()
    paragraph.add_run().add_break(WD_BREAK.PAGE)


def configure_document(doc: Document) -> None:
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.65)
    section.left_margin = Inches(0.7)
    section.right_margin = Inches(0.7)
    section.header_distance = Inches(0.32)
    section.footer_distance = Inches(0.32)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Aptos"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Aptos")
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = BLACK
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.12

    body = styles["Body Text"]
    body.font.name = "Aptos"
    body._element.rPr.rFonts.set(qn("w:eastAsia"), "Aptos")
    body.font.size = Pt(10.5)
    body.font.color.rgb = BLACK
    body.paragraph_format.space_after = Pt(7)
    body.paragraph_format.line_spacing = 1.14

    title = styles["Title"]
    title.font.name = "Aptos Display"
    title._element.rPr.rFonts.set(qn("w:eastAsia"), "Aptos Display")
    title.font.size = Pt(28)
    title.font.bold = True
    title.font.color.rgb = BLACK
    title.paragraph_format.space_after = Pt(12)

    subtitle = styles["Subtitle"]
    subtitle.font.name = "Aptos"
    subtitle._element.rPr.rFonts.set(qn("w:eastAsia"), "Aptos")
    subtitle.font.size = Pt(14)
    subtitle.font.color.rgb = MUTED
    subtitle.paragraph_format.space_after = Pt(18)

    for name, size, before, after in (
        ("Heading 1", 18, 12, 7),
        ("Heading 2", 13, 10, 5),
        ("Heading 3", 11, 8, 4),
    ):
        style = styles[name]
        style.font.name = "Aptos Display"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Aptos Display")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = BLACK
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    for list_style in ("List Bullet", "List Number"):
        style = styles[list_style]
        style.font.name = "Aptos"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Aptos")
        style.font.size = Pt(10.5)
        style.font.color.rgb = BLACK
        style.paragraph_format.left_indent = Inches(0.28)
        style.paragraph_format.first_line_indent = Inches(-0.18)
        style.paragraph_format.line_spacing = 1.1

    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.LEFT
    header_run = header.add_run("OnlyMyPDF Phase 1 Implementation Checkpoint")
    header_run.font.name = "Aptos"
    header_run.font.size = Pt(8.5)
    header_run.font.color.rgb = MUTED

    footer = section.footer.paragraphs[0]
    add_page_number(footer)

    doc.core_properties.title = "OnlyMyPDF Phase 1 Dependable Beta Implementation Checkpoint"
    doc.core_properties.subject = "Evidence based implementation and readiness checkpoint"
    doc.core_properties.author = "OnlyMyPDF Development"
    doc.core_properties.comments = "Generated from verified Phase 1 implementation evidence"


def build_document() -> Document:
    doc = Document()
    configure_document(doc)

    cover_spacer = doc.add_paragraph()
    cover_spacer.paragraph_format.space_after = Pt(52)
    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    title.add_run("OnlyMyPDF Phase 1 Dependable Beta Implementation Checkpoint")
    clear_paragraph_borders(title)
    subtitle = doc.add_paragraph(style="Subtitle")
    subtitle.add_run("Developer evidence report and live activation gate")

    meta = doc.add_table(rows=4, cols=2)
    meta.alignment = WD_TABLE_ALIGNMENT.LEFT
    meta.autofit = False
    set_table_borders(meta, color="FFFFFF", size="0")
    meta_rows = [
        ("Report date", "17 September 2026"),
        ("Branch", "phase1-dependable-beta"),
        ("Current product readiness", "80 out of 100"),
        ("Phase 1 completion", "88 percent"),
    ]
    for index, (label, value) in enumerate(meta_rows):
        meta.rows[index].cells[0].width = Inches(1.9)
        meta.rows[index].cells[1].width = Inches(4.7)
        for cell in meta.rows[index].cells:
            set_cell_margins(cell, top=75, start=0, bottom=75, end=120)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p1 = meta.rows[index].cells[0].paragraphs[0]
        p1.paragraph_format.space_after = Pt(0)
        r1 = p1.add_run(label)
        r1.bold = True
        r1.font.size = Pt(10.5)
        p2 = meta.rows[index].cells[1].paragraphs[0]
        p2.paragraph_format.space_after = Pt(0)
        r2 = p2.add_run(value)
        r2.font.size = Pt(10.5)

    cover_note = doc.add_paragraph(style="Body Text")
    cover_note.paragraph_format.space_before = Pt(34)
    cover_note.add_run("Main conclusion  ").bold = True
    cover_note.add_run(
        "Phase 1 application code and local verification are complete. Live infrastructure activation and production evidence remain required before Phase 1 can be marked complete and before Phase 2 begins."
    )
    cover_scope = doc.add_paragraph(style="Body Text")
    cover_scope.add_run("Scope  ").bold = True
    cover_scope.add_run(
        "Conversion validity, quality corpus, queue durability, worker recovery, private storage, retention cleanup, upload boundaries, operational telemetry, and release gates."
    )

    add_page_break(doc)
    add_heading(doc, "Decision and current status", 1)
    add_body(
        doc,
        "Phase 1 is locally implemented and verified at 88 percent. It cannot be described as 100 percent complete because the target Supabase project, Upstash queue, isolated worker host, scheduled cleanup, external monitoring, and real CDN upload path are not configured in the current environment.",
    )
    add_body(
        doc,
        "The correct release decision is to hold Phase 2. Complete the six production activation checks in this report, collect the resulting evidence, and then update the Phase 1 status to 100 percent.",
    )
    add_table(
        doc,
        ["Measure", "Starting point", "Current", "After live gate"],
        [
            ["Product readiness", "70 out of 100", "80 out of 100", "82 out of 100"],
            ["Phase 1 completion", "0 percent", "88 percent", "100 percent"],
            ["Phase 2 status", "Not started", "Not started", "Eligible to start"],
        ],
        [2.1, 1.65, 1.65, 1.65],
        font_size=9.5,
        center_columns={1, 2, 3},
    )
    add_heading(doc, "Environment status", 2)
    add_body(
        doc,
        "Local engine paths and the cron authentication secret are available. Production data services are absent, so local success cannot prove storage privacy, durable queue operation, scheduled deletion, or alert delivery in production.",
    )
    add_table(
        doc,
        ["Configuration", "Status", "Meaning"],
        [
            ["CRON SECRET", "Available", "Local authenticated worker and cleanup routes can run"],
            ["PDF2DOCX PYTHON", "Available", "Local PDF to Word engine can run"],
            ["LIBREOFFICE PATH", "Available", "Local Office conversion engine can run"],
            ["Supabase URL and service role", "Missing", "Migration and private storage are not live verified"],
            ["Upstash Redis URL and token", "Missing", "Persistent production queue is not live verified"],
            ["HEALTH CHECK SECRET", "Missing", "Detailed production health endpoint is not activated"],
            ["ConvertAPI secret", "Missing", "Paid fallback is not configured or tested"],
        ],
        [2.35, 1.25, 3.45],
        font_size=9.0,
        center_columns={1},
    )

    add_page_break(doc)
    add_heading(doc, "Issues found and resolved", 1)
    add_body(
        doc,
        "Fourteen reliability gaps were identified during Phase 1. Each item below is implemented and verified locally. The production qualification column states where a live environment is still required.",
    )
    resolved_rows = [
        ["1", "API restart could lose background work", "Persistent FIFO queue, private staged input, and queued running done error lifecycle", "Live Redis and storage check pending"],
        ["2", "Worker crash had no recovery path", "Processing list, 15 minute lease recovery, isolated worker command, and constrained worker container", "Deploy and force one recovery scenario"],
        ["3", "A corrupt nonempty output could be marked complete", "Format aware openability, signature, page, slide, sheet, and package checks", "Resolved in code and tests"],
        ["4", "First guest job could fail polling with Access denied", "New guest cookie is forwarded into the same request and response", "Resolved by direct localhost job"],
        ["5", "Conversion outcome and latency were not measurable", "Event telemetry records engine, fallback, queue time, duration, validity, attempts, and failures", "Apply migration and collect live history"],
        ["6", "Crashed staged input could survive Redis expiry", "Hourly storage scan removes PDF to Word inputs older than two hours and records failures", "Observe scheduled cleanup live"],
        ["7", "Signed URL could outlive file retention", "URL lifetime is capped by two hours and the database expiry deadline; bucket migration forces private access", "Apply migration and inspect live object policy"],
        ["8", "Declared upload size could differ from returned bytes", "Server rejects size mismatch and enforces exact plan boundaries", "Application validator resolved; real CDN path pending"],
        ["9", "No declared conversion quality corpus", "Deterministic 200 document corpus covers text, tables, scans, Hindi images, orientation, forms, fonts, and larger files", "Resolved for the tested operations"],
        ["10", "Large files crossed the app request body", "Owner bound signed upload sends configured clients directly to private Supabase storage", "Run real 25 and 200 MB production uploads"],
        ["11", "Queued PDF passwords needed safe worker transport", "AES 256 GCM encrypted job payload; secret removed after completion or failure", "Resolved in code and tests"],
        ["12", "A deployed worker could be silently dead", "Upstash heartbeat every 15 seconds; health fails after 60 seconds without a healthy worker", "Deploy worker and alert on failure"],
        ["13", "Nested abandoned direct uploads escaped cleanup", "Bounded paginated directory traversal removes expired uploads UUID input PDF objects", "Observe scheduled cleanup live"],
        ["14", "Output storage failure could still mark a distributed job done", "Production and Redis jobs fail closed until validated DOCX is in private storage", "Resolved in code; live failure drill pending"],
    ]
    add_table(
        doc,
        ["ID", "Problem", "Implemented resolution", "Qualification"],
        resolved_rows,
        [0.45, 2.0, 3.1, 1.5],
        font_size=7.7,
        center_columns={0},
    )

    add_page_break(doc)
    add_heading(doc, "Verification evidence", 1)
    add_body(
        doc,
        "The following results come from actual generated files, opened output packages, rendered comparisons, direct localhost API use, automated tests, type checking, linting, a production build, and a production dependency audit.",
    )
    evidence_rows = [
        ["Declared corpus", "Pass", "200 of 200 jobs succeeded against a 99.5 percent threshold"],
        ["Extractable text", "Pass", "155 of 155 applicable marker checks passed"],
        ["Rendered fidelity", "Pass", "40 of 40 comparisons passed; minimum similarity 97.99 percent"],
        ["Simple latency", "Pass", "Corpus p95 was 309 ms, below 15 seconds"],
        ["Office engines", "Pass", "7 of 7 PDF and Office conversions opened successfully"],
        ["PDF to Word", "Pass", "Word COM; 14.353 s; valid DOCX; 1,330 extracted characters"],
        ["PDF to Excel", "Pass", "2.126 s; valid workbook; two worksheets"],
        ["PDF to PowerPoint text", "Pass", "3.617 s; two slides; 1,330 characters; two editable slides reported"],
        ["PDF to PowerPoint scan", "Pass with limit", "3.353 s; valid two slide image presentation; no text expected"],
        ["Office to PDF", "Pass", "Word 3.313 s; Excel 3.678 s; PowerPoint 5.989 s"],
        ["Final localhost API job", "Pass", "1,096,734 byte PDF to 66,511 byte DOCX; 48 ms queue; 6.301 s processing"],
        ["Artifact validation", "Pass", "DOCX signature and openability valid; 19 entries; one page; 1,052 text characters"],
        ["Protected download", "Pass", "First valid DOCX returned; repeated consume returned 404"],
        ["Upload limits", "Pass locally", "Exact 25 and 200 MB accepted; one byte over each limit rejected"],
        ["New upload security cleanup tests", "Pass", "11 of 11"],
        ["Regression suite", "Pass", "605 of 605 across 124 files"],
        ["Type checking", "Pass", "TypeScript completed with no errors"],
        ["Lint", "Pass", "Zero errors and 18 existing warnings"],
        ["Production build", "Pass", "154 of 154 static pages; upload and worker routes generated"],
        ["Production dependencies", "Pass", "Zero moderate high or critical production vulnerabilities"],
    ]
    add_table(
        doc,
        ["Gate", "Result", "Measured evidence"],
        evidence_rows,
        [2.0, 1.2, 4.05],
        font_size=8.7,
        center_columns={1},
    )

    add_heading(doc, "Phase 1 workstream status", 1)
    add_body(
        doc,
        "The completion score separates local implementation from production activation. This prevents a code ready feature from being reported as production verified before its infrastructure exists.",
    )
    workstream_rows = [
        ["200 document quality corpus", "20 percent", "20 percent", "Local gate passed"],
        ["Output validation and fidelity evidence", "15 percent", "15 percent", "Local gate passed"],
        ["Engine routing and fallback evidence", "10 percent", "10 percent", "Local gate passed"],
        ["Queue worker retry and crash recovery", "20 percent", "17 percent", "Production worker pending"],
        ["Private storage and deletion evidence", "15 percent", "12 percent", "Live retention observation pending"],
        ["25 and 200 MB upload path", "10 percent", "7 percent", "Direct storage ready; live proof pending"],
        ["Monitoring and alerting", "10 percent", "7 percent", "External monitor and live history pending"],
        ["Total", "100 percent", "88 percent", "Do not start Phase 2"],
    ]
    add_table(
        doc,
        ["Workstream", "Weight", "Complete", "Status"],
        workstream_rows,
        [3.0, 1.1, 1.1, 2.05],
        font_size=9.0,
        center_columns={1, 2},
    )
    add_heading(doc, "Implemented processing path", 2)
    add_body(
        doc,
        "The PDF to Word path now follows a durable job lifecycle. Configured browsers upload large files directly to private storage with an owner bound grant. A queue entry is created, an isolated worker heartbeat proves the worker is alive, the worker claims and validates the input, conversion attempts and fallback are recorded, the output must persist to private storage, and the protected download is consumed once. Cleanup recursively removes expired staged input.",
    )
    add_table(
        doc,
        ["Stage", "Responsibility", "Recorded evidence"],
        [
            ["Request", "Validate plan limit owner and declared input bytes", "Owner and input size"],
            ["Direct upload", "Issue an expiring owner bound private storage grant", "Path size owner and one time claim"],
            ["Queue", "Persist FIFO job and staged private input", "Queued time and depth"],
            ["Worker", "Claim recover retry and route engines", "Attempts engines fallback and duration"],
            ["Validation", "Open and inspect final artifact", "Validity type counts and text"],
            ["Delivery", "Authorize and consume protected result", "Downloaded once then unavailable"],
            ["Cleanup", "Delete expired rows and staged objects", "Deleted and failed counts"],
        ],
        [1.2, 3.35, 2.7],
        font_size=9.0,
    )

    add_page_break(doc)
    add_heading(doc, "Required live activation gate", 1)
    add_body(
        doc,
        "Complete all six actions in the target environment. Each action must leave reviewable evidence. Configuration alone is not sufficient.",
    )
    live_items = [
        "Configure the target Supabase and Upstash credentials, apply database migrations 001 through 021, and confirm that the pdf files bucket is private.",
        "Deploy the dedicated worker from Dockerfile.worker or run npm run worker:conversions on an isolated host with restart policy and resource limits.",
        "Call the authenticated health endpoint with HEALTH CHECK SECRET and confirm private storage, fresh worker heartbeat, dedicated worker mode, queue depth, output validity, conversion latency, fallback rate, and cleanup status are healthy.",
        "Observe one successful scheduled cleanup cycle and one controlled failed object retry. Confirm that no staged input survives beyond the two hour limit.",
        "Send real 25 MB Free and 200 MB Pro files through CDN or proxy, application, queue, worker, private storage, and final download. Record total time, peak resource use, output validity, and downloaded file size.",
        "Connect an external monitor to detailed health degradation and a worker or container restart alert. Trigger each alert once and record delivery evidence.",
    ]
    add_numbered(doc, live_items)
    add_heading(doc, "Acceptance evidence", 2)
    add_table(
        doc,
        ["Check", "Required proof", "Completion rule"],
        [
            ["Migration and storage", "Migration history and private bucket policy", "No public object read is possible"],
            ["Queue and worker", "Worker log plus queued running done lifecycle", "Restart does not lose a job"],
            ["Cleanup", "Cleanup run record and empty expired object result", "Expired inputs are removed on schedule"],
            ["25 MB Free path", "End to end timing and valid downloaded artifact", "Plan limit and output contract pass"],
            ["200 MB Pro path", "End to end timing resource use and valid artifact", "No proxy storage or worker truncation"],
            ["Monitoring", "Delivered degradation and restart alerts", "A responsible operator receives both"],
        ],
        [1.65, 3.65, 1.95],
        font_size=9.0,
    )

    add_page_break(doc)
    add_heading(doc, "Files and reproducible evidence", 1)
    add_body(
        doc,
        "The repository keeps the generators, reports, migration, tests, and operational documentation needed to reproduce this checkpoint.",
    )
    add_table(
        doc,
        ["Purpose", "Repository path"],
        [
            ["Phase 1 checkpoint", "docs/PHASE_1_IMPLEMENTATION_CHECKPOINT.md"],
            ["Corpus manifest", "quality/phase1-corpus/manifest.json"],
            ["Corpus result", "quality/phase1-corpus/latest-report.json"],
            ["Office engine result", "quality/phase1-corpus/engine-report.json"],
            ["Upload boundary result", "quality/phase1-corpus/upload-boundary-report.json"],
            ["Final local API smoke", "quality/phase1-corpus/local-api-smoke-report.json"],
            ["Operational migration", "supabase/migrations/021_phase1_conversion_operations.sql"],
            ["Queue worker", "workers/conversion-worker.ts"],
            ["Worker container", "Dockerfile.worker"],
            ["Output validator", "src/lib/services/conversion-output-validation.ts"],
            ["Telemetry", "src/lib/ops/conversion-telemetry.ts"],
            ["Health endpoint", "src/app/api/health/route.ts"],
            ["Operations guide", "docs/OPERATIONS.md"],
        ],
        [2.35, 4.9],
        font_size=9.2,
    )
    add_heading(doc, "Final phase decision", 2)
    add_body(
        doc,
        "Phase 1 local implementation is complete. The current checkpoint is 88 percent because production infrastructure proof is still absent. Do not begin Phase 2 until all six live activation checks pass and the resulting evidence is added to the repository.",
    )
    add_body(
        doc,
        "When the live gate passes, update the Phase 1 score to 100 percent, product readiness to 82 out of 100, and record the production date, environment, file sizes, processing times, output validity, cleanup results, and alert delivery proof.",
    )
    add_heading(doc, "Known limits", 2)
    add_bullets(
        doc,
        [
            "The corpus proves the tested operations and fixtures. It does not prove that every real world PDF can be converted with identical layout.",
            "The 200 MB result currently proves the application validation boundary. It does not yet prove the complete production network and worker path.",
            "The image only scan presentation is valid by design and contains no extractable text. OCR quality is a separate product capability.",
            "Paid ConvertAPI fallback remains unconfigured and unverified.",
        ],
    )

    return doc


def main() -> None:
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    document = build_document()
    document.save(OUTPUT)
    print(f"Created {OUTPUT}")
    print(f"Bytes {OUTPUT.stat().st_size}")


if __name__ == "__main__":
    main()
