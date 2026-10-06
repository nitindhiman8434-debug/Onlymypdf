"""Readable editable fallback for OCR transcripts containing only Latin letters.

This script check is deliberately not language identification. Text and spacing
come from actual OCR; no expected words, inferred corrections or images are added.
"""

from __future__ import annotations

import unicodedata
import os
import math
from collections import Counter
from pathlib import Path
import shutil
import subprocess
import tempfile


class EmptyOcrPageError(ValueError):
    def __init__(self, page_number: int):
        self.page_number = page_number
        super().__init__(f"OCR did not recognize readable text on page {page_number}")


class OcrRuntimeUnavailableError(RuntimeError):
    """The executable or its requested trained data cannot be used."""


def recognize_ocr_transcripts(source_pdf: str, *, language: str, dpi: int,
                              page_layouts: list[dict | None] | None = None) -> list[str]:
    """Read Tesseract's own word boundaries rather than reconstructing PDF glyphs.

    Called only for the Latin-script fallback selected by the converter. One
    bounded PSM 3 pass per original page keeps actual lines, spaces and repeated
    content. Optional TSV geometry is emitted by that same pass; incomplete or
    ambiguous table evidence leaves the original transcript unchanged.
    """
    import pymupdf as fitz

    executable = shutil.which(os.environ.get("TESSERACT_PATH", "tesseract"))
    if executable is None:
        raise OcrRuntimeUnavailableError("Tesseract executable unavailable")

    transcripts: list[str] = []
    if page_layouts is not None:
        page_layouts.clear()
    with fitz.open(source_pdf) as source, tempfile.TemporaryDirectory(prefix="pdf2docx-ocr-text-") as temp:
        image_path = Path(temp) / "page.png"
        output_base = Path(temp) / "recognized"
        output_text = output_base.with_suffix(".txt")
        output_tsv = output_base.with_suffix(".tsv")
        for page in source:
            pixmap = page.get_pixmap(dpi=dpi, colorspace=fitz.csRGB, alpha=False)
            prepared_image = None
            if page_layouts is not None:
                from ocr_ruled_tables import prepare_ruled_page_ocr_image
                prepared_image = prepare_ruled_page_ocr_image(pixmap)
            if prepared_image is None:
                pixmap.save(image_path)
            else:
                image_path.write_bytes(prepared_image)
            # A successful prior page must never stand in for a missing result.
            output_text.unlink(missing_ok=True)
            output_tsv.unlink(missing_ok=True)
            formats = ["txt", "tsv"] if page_layouts is not None else ["txt"]
            try:
                result = subprocess.run(
                    [executable, str(image_path), str(output_base), "-l", language,
                     "--dpi", str(dpi), "--psm", "3", *formats],
                    capture_output=True, text=True, timeout=90, check=False,
                )
            except FileNotFoundError as exc:
                raise OcrRuntimeUnavailableError("Tesseract executable unavailable") from exc
            if result.returncode != 0:
                diagnostics = (result.stderr or "").lower()
                if any(marker in diagnostics for marker in (
                    "error opening data file", "failed loading language", "couldn't load any languages",
                )):
                    raise OcrRuntimeUnavailableError("Tesseract requested language data unavailable")
                raise RuntimeError(f"Tesseract text recognition failed (exit {result.returncode})")
            if not output_text.is_file():
                raise RuntimeError("Tesseract text recognition produced no transcript")
            # Keep actual blank lines, repeated rows and intra-line spacing.
            transcript = output_text.read_text(encoding="utf-8").replace("\r\n", "\n").replace("\r", "\n")
            transcripts.append(transcript)
            if page_layouts is not None:
                layout = None
                if output_tsv.is_file() and supports_latin_ocr([transcript]):
                    from ocr_ruled_tables import build_ruled_page_layout
                    try:
                        tsv_text = output_tsv.read_text(encoding="utf-8")
                    except UnicodeError:
                        tsv_text = ""  # Invalid optional geometry must not discard valid TXT.
                    layout = build_ruled_page_layout(pixmap, tsv_text, transcript)
                page_layouts.append(layout)
    return transcripts


def supports_latin_ocr(transcripts: list[str]) -> bool:
    """Select Latin-script text only; empty pages are validated separately."""
    letters = [character for text in transcripts for character in text if character.isalpha()]
    return bool(letters) and all("LATIN" in unicodedata.name(character, "") for character in letters)


def validate_page_transcripts(transcripts: list[str], page_count: int) -> None:
    if len(transcripts) != page_count:
        raise ValueError("OCR transcript count does not match the source page count")
    for page_number, text in enumerate(transcripts, 1):
        if not text.strip():
            raise EmptyOcrPageError(page_number)


def _validated_layout_blocks(layout: dict | None, transcript: str) -> list[dict] | None:
    """Validate all blocks before writing; keep the earlier one-table contract."""
    if layout is None:
        return None
    blocks = layout.get("blocks") if "blocks" in layout else [
        {"kind": "paragraphs", "lines": layout.get("before")},
        {"kind": "table", "rows": layout.get("rows"), "columnWidths": layout.get("columnWidths")},
        {"kind": "paragraphs", "lines": layout.get("after")},
    ]
    if not isinstance(blocks, list) or not blocks:
        raise ValueError("OCR table layout requires ordered blocks")
    texts, tables = [], 0
    for block in blocks:
        if not isinstance(block, dict):
            raise ValueError("Invalid OCR layout block")
        if block.get("kind") == "paragraphs":
            lines = block.get("lines")
            if not isinstance(lines, list) or not all(isinstance(line, str) for line in lines):
                raise ValueError("Invalid OCR paragraph block")
            texts.extend(lines)
        elif block.get("kind") == "table":
            rows, widths = block.get("rows"), block.get("columnWidths")
            if (not isinstance(rows, list) or not 2 <= len(rows) <= 40
                    or not isinstance(widths, list) or not 2 <= len(widths) <= 8
                    or any(not isinstance(row, list) or len(row) != len(widths)
                           or not all(isinstance(cell, str) for cell in row) for row in rows)
                    or any(not isinstance(width, (float, int)) or not math.isfinite(width)
                           or width <= 0 for width in widths)
                    or not math.isclose(sum(widths), 1, abs_tol=1e-6)):
                raise ValueError("Invalid OCR table matrix or column widths")
            texts.extend(cell for row in rows for cell in row)
            tables += 1
        else:
            raise ValueError("Unknown OCR layout block kind")
    if not 1 <= tables <= 2:
        raise ValueError("OCR layout must contain one or two tables")
    if Counter(" ".join(texts).split()) != Counter(transcript.split()):
        raise ValueError("OCR layout must preserve every recognized word")
    return blocks


def write_english_ocr_docx(source_pdf: str, output_path: str, transcripts: list[str], *,
                           page_layouts: list[dict | None] | None = None) -> None:
    """Write full-width OCR lines and preserve each source page's section size.

    Dimensions must come from the original input PDF, not a generated OCR PDF
    whose image resolution may change its physical page size.
    Sections start on a new page. Dense text can reflow onto more than one Word
    page; this does not promise a fixed rendered page count or original typography.
    """
    import pymupdf as fitz
    from docx import Document
    from docx.enum.section import WD_ORIENT, WD_SECTION_START
    from docx.oxml import OxmlElement
    from docx.shared import Pt

    with fitz.open(source_pdf) as source:
        validate_page_transcripts(transcripts, len(source))
        if not supports_latin_ocr(transcripts):
            raise ValueError("The readable English OCR fallback requires Latin-script text")
        if page_layouts and len(page_layouts) != len(transcripts):
            raise ValueError("OCR table layout count does not match the source page count")
        blocks_by_page = [_validated_layout_blocks(layout, transcript)
                          for layout, transcript in zip(page_layouts, transcripts)] if page_layouts else [None] * len(transcripts)

        word = Document()
        normal = word.styles["Normal"]
        normal.font.name = "Arial"
        normal.font.size = Pt(11)
        normal.paragraph_format.space_before = Pt(0)
        normal.paragraph_format.space_after = Pt(4)
        normal.paragraph_format.line_spacing = 1.15

        for page_index, page in enumerate(source):
            section = word.sections[0] if page_index == 0 else word.add_section(WD_SECTION_START.NEW_PAGE)
            width, height = page.rect.width, page.rect.height
            section.orientation = WD_ORIENT.LANDSCAPE if width > height else WD_ORIENT.PORTRAIT
            section.page_width, section.page_height = Pt(width), Pt(height)
            section.left_margin = section.right_margin = Pt(min(43.2, width / 10))
            section.top_margin = section.bottom_margin = Pt(min(43.2, height / 10))

            def add_lines(lines):
                for line in lines:
                    paragraph = word.add_paragraph(line)
                    paragraph.paragraph_format.keep_with_next = False
                    paragraph.paragraph_format.widow_control = True

            def add_table(block):
                rows = block["rows"]
                table = word.add_table(rows=len(rows), cols=len(rows[0]))
                table.style = "Table Grid"
                table.autofit = False
                usable_width = section.page_width.pt - section.left_margin.pt - section.right_margin.pt
                for index, fraction in enumerate(block["columnWidths"]):
                    table.columns[index].width = Pt(usable_width * fraction)
                    for cell in table.columns[index].cells:
                        cell.width = Pt(usable_width * fraction)
                for row_index, values in enumerate(rows):
                    row = table.rows[row_index]
                    row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
                    for cell, value in zip(row.cells, values):
                        cell.text = value
                        for paragraph in cell.paragraphs:
                            paragraph.paragraph_format.space_after = Pt(2)
                            paragraph.paragraph_format.keep_with_next = False

            blocks = blocks_by_page[page_index]
            if blocks:
                previous_was_table = False
                for block in blocks:
                    if block["kind"] == "paragraphs":
                        add_lines(block["lines"])
                        if block["lines"]:
                            previous_was_table = False
                    else:
                        # Word can coalesce consecutive tables without a separator.
                        if previous_was_table:
                            word.add_paragraph()
                        add_table(block)
                        previous_was_table = True
                continue

            # Keep the exact original lines when no confident ruled table exists.
            for line in transcripts[page_index].splitlines():
                paragraph = word.add_paragraph(line)
                paragraph.paragraph_format.keep_with_next = False
                paragraph.paragraph_format.widow_control = True

        word.save(output_path)
